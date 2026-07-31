"""
PO_load.py
----------
Loads Purchase Order headers and lines from all .xlsx files in the
configured PurchaseOrder folder into QAD via the purchaseOrderHeaders /
purchaseOrderLines APIs.

Behaviour
---------
- Reads base_url and auth from config.json (via config.py)
- Fetches one OAuth token per run; refreshes automatically on 401
  (unchanged token pattern: api_get()/api_post() retry once internally)
- Skips rows with Status = DONE
- Error/response handling now shares the same framework as
  Customer_load.py:
    * qad_response_utils.normalize_response() turns every QAD POST
      response into a uniform (success, errors[]) shape, regardless of
      whether QAD sent a submitResult-wrapped envelope or a bare
      top-level errors[] (gateway/permission failures).
    * excel_format_utils.resolve_qad_errors() maps each errors[]
      fieldName back to the Excel column name via PO_HEADER_FIELD_TO_COLUMN
      / PO_LINE_FIELD_TO_COLUMN below, so the exact bad cell gets
      highlighted instead of just col 1 + Error.
    * excel_format_utils.mark_done() / mark_error() replace the old
      mark_row_success() / mark_row_error() helpers and share the same
      RED_FILL / CLEAR_FILL constants used by every other loader.
- Lightweight mandatory-field check runs before any API call, for both
  Header rows and Line rows.
- Network errors, persistent-401 (token refresh failure), and gateway/
  permission errors are all surfaced as normalized error dicts so they
  flow through the exact same resolve_qad_errors() / mark_error() path
  as ordinary field validation errors.
- Saves workbook after every row so progress survives a crash.
- Renames file to error_<name> if any failures occurred (unchanged).
- Returns (total_success, total_fail)

Sheet layout (UNCHANGED)
------------------------
Workbook must contain two sheets: 'Header' and 'Lines'.
Each sheet must have a header row (row 1) with column names.

Purchase Order business logic / API workflow — UNCHANGED
----------------------------------------------------------
  Header : single POST to purchaseOrderHeaders
  Lines  : GET  init (QAD assigns line number)
           POST fieldChange(siteCode)
           POST fieldChange(itemCode)
           GET  isReceived check (fire-and-forget)
           POST fieldChange(quantityOrdered)
           POST fieldChange(purchaseCost)
           POST sync to purchaseOrderLinesGrid
None of the endpoints, payload fields, sequencing, or line-number
handling have been changed — only how responses are parsed and how
Excel is annotated.
"""

import os
import sys
import time
import requests
import openpyxl
from datetime import datetime
from collections import defaultdict

# ── Progress callback (used by main.py SSE streaming) ──────────────────

_progress_callback = None

def set_progress_callback(callback):
    global _progress_callback
    _progress_callback = callback


# ── Path / config ─────────────────────────────────────────────────────────
ROOT_DIR    = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT_DIR)
sys.path.insert(0, SCRIPTS_DIR)
from config import CONFIG
from excel_format_utils import RED_FILL, CLEAR_FILL, mark_done, mark_error, resolve_qad_errors
from qad_response_utils import normalize_response

_BASE_URL = CONFIG["qad"]["base_url"]

# ── API endpoint templates (UNCHANGED) ─────────────────────────────────────
HEADER_URL       = f"{_BASE_URL}/api/erp/purchaseOrderHeaders?viewUri=urn:be:com.qad.purchasing.purchaseorders.IPurchaseOrderHeader"
INIT_LINE_URL    = f"{_BASE_URL}/api/erp/purchaseOrderLinesGrid?initialize=true&domainCode={{domain}}&purchaseOrderNumber={{po}}"
FIELD_CHANGE_URL = f"{_BASE_URL}/api/erp/purchaseOrderLines/fieldChangeV2?fieldName={{fieldName}}"
IS_RECEIVED_URL  = f"{_BASE_URL}/api/erp/purchaseOrderLines/isReceivedPurchaseOrderLineV2?domainCode={{domain}}&purchaseOrderNumber={{po}}&purchaseOrderLine={{line}}"
SYNC_LINE_URL    = f"{_BASE_URL}/api/erp/purchaseOrderLinesGrid"

# ── Mandatory columns — Header sheet ───────────────────────────────────────
MANDATORY_HEADER_COLUMNS = [
    "PO Number",
    "Domain Code",
    "Supplier Code",
    "Currency",
    "Credit Terms",
    "Daybook Set",
    "Ship To Site",
    "Bill To Code",
]

# ── Mandatory columns — Lines sheet ─────────────────────────────────────────
MANDATORY_LINE_COLUMNS = [
    "PO Number",
    "Site Code",
    "Item Code",
    "Quantity Ordered",
    "Unit Price",
]

# ── QAD API fieldName → Excel column name (Header / purchaseOrderHeaders) ──
# Extend this map whenever a new QAD fieldName is observed in errors[]
# that should highlight a specific Excel cell. Anything NOT in this map
# still shows up in the Error message text — it just won't get its own
# highlighted cell.
PO_HEADER_FIELD_TO_COLUMN = {
    "purchaseOrderNumber": "PO Number",
    "domainCode":          "Domain Code",
    "supplierCode":        "Supplier Code",
    "currencyCode":        "Currency",
    "creditTermsCode":     "Credit Terms",
    "daybookSetCode":      "Daybook Set",
    "shipToSite":          "Ship To Site",
    "billToCode":          "Bill To Code",
    "taxEnvironment":      "Tax Environment",
    "languageCode":        "Language Code",
    "contact":             "Contact",
    "remarks":             "Remarks",
    "orderDate":           "Order Date",
    "dueDate":             "Due Date",
}

# ── QAD API fieldName → Excel column name (Lines / purchaseOrderLines) ─────
PO_LINE_FIELD_TO_COLUMN = {
    "purchaseOrderNumber": "PO Number",
    "siteCode":            "Site Code",
    "itemCode":            "Item Code",
    "quantityOrdered":     "Quantity Ordered",
    "purchaseCost":        "Unit Price",
    "dueDate":             "Due Date",
}


# =============================================================================
# 1. AUTHENTICATION  (unchanged token pattern)
# =============================================================================

def _fetch_token() -> str:
    token_url = f"{_BASE_URL}/oauth/token"
    resp = requests.post(token_url, params=CONFIG["qad"]["auth"], timeout=30)
    resp.raise_for_status()
    token = resp.json().get("access_token")
    if not token:
        raise RuntimeError("OAuth response did not contain access_token")
    return token


class TokenManager:
    """Holds one token for the run; refreshes on demand (401)."""

    def __init__(self):
        self._token: str | None = None

    def get(self) -> str:
        if self._token is None:
            self._token = _fetch_token()
        return self._token

    def refresh(self) -> str:
        self._token = _fetch_token()
        return self._token

    def headers(self) -> dict:
        return {
            "Content-Type":  "application/json",
            "Authorization": f"Bearer {self.get()}",
        }


# =============================================================================
# 2. HTTP HELPERS  (auto-refresh on 401, one retry — UNCHANGED behaviour)
# =============================================================================

def api_get(url: str, tm: TokenManager) -> tuple[requests.Response, dict]:
    """GET with one automatic token-refresh retry on 401."""
    for attempt in range(2):
        resp = requests.get(url, headers=tm.headers(), timeout=30)
        if resp.status_code == 401 and attempt == 0:
            tm.refresh()
            continue
        try:
            return resp, resp.json()
        except Exception:
            return resp, {}
    return resp, {}


def api_post(url: str, payload: dict, tm: TokenManager) -> tuple[requests.Response, dict]:
    """POST with one automatic token-refresh retry on 401."""
    for attempt in range(2):
        resp = requests.post(url, json=payload, headers=tm.headers(), timeout=30)
        if resp.status_code == 401 and attempt == 0:
            tm.refresh()
            continue
        try:
            return resp, resp.json()
        except Exception:
            return resp, {}
    return resp, {}


# ── Normalized-error wrappers ────────────────────────────────────────────
# These wrap api_get()/api_post() so every caller (header creation, line
# creation) gets back the same (resp, resp_json, errors) shape used
# throughout Customer_load.py — errors is either None (success) or a
# normalized list of {"fieldName", "message", "fieldValue", "code"} dicts,
# ready to be handed straight to resolve_qad_errors().

def _network_error(exc: Exception) -> dict:
    return {"fieldName": None, "message": f"Network error: {exc}", "fieldValue": None, "code": None}


def _token_error() -> dict:
    return {"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}


def _http_error(status_code: int) -> dict:
    return {"fieldName": None, "message": f"Unexpected response — HTTP {status_code}", "fieldValue": None, "code": None}


def safe_get(url: str, tm: TokenManager) -> tuple[requests.Response | None, dict | None, list | None]:
    """GET wrapped with network-error trapping. Does NOT run normalize_response
    (GET responses here are plain data payloads, not submit envelopes) —
    mirrors get_customer()/get_mfg_customer() in Customer_load.py, which are
    also not normalized."""
    try:
        resp, resp_json = api_get(url, tm)
    except requests.RequestException as e:
        return None, None, [_network_error(e)]

    if resp.status_code == 401:
        return resp, resp_json, [_token_error()]
    if resp.status_code != 200:
        return resp, resp_json, [_http_error(resp.status_code)]
    return resp, resp_json, None


def safe_post(url: str, payload: dict, tm: TokenManager) -> tuple[requests.Response | None, dict | None, list | None]:
    """POST wrapped with network-error trapping and normalize_response(),
    exactly like post_customer()/post_mfg_customer() in Customer_load.py."""
    try:
        resp, resp_json = api_post(url, payload, tm)
    except requests.RequestException as e:
        return None, None, [_network_error(e)]

    if resp.status_code == 401:
        return resp, resp_json, [_token_error()]

    success, errors = normalize_response(resp_json, resp.status_code)
    if not success:
        if not errors:
            errors = [_http_error(resp.status_code)]
        return resp, resp_json, errors
    return resp, resp_json, None


# =============================================================================
# 3. VALUE EXTRACTORS  (unchanged)
# =============================================================================

def sv(row_data: dict, key: str, default: str = "") -> str:
    val = row_data.get(key, default)
    return str(val).strip() if val is not None else default


def fv(row_data: dict, key: str, default: float = 0.0) -> float:
    try:
        return float(row_data.get(key, default))
    except (TypeError, ValueError):
        return default


def to_iso_date(val) -> str:
    """Convert a cell value or datetime to ISO 8601 UTC string."""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%dT%H:%M:%S.000Z")
    if val and isinstance(val, str):
        try:
            return datetime.strptime(val, "%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%dT%H:%M:%S.000Z")
        except ValueError:
            return val
    return val or ""


# =============================================================================
# 4. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict, columns: list[str]) -> list[str]:
    missing = []
    for col in columns:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 5. WORKBOOK HELPERS  (sheet/header layout UNCHANGED — row 1 = headers)
# =============================================================================

def ensure_columns(ws, *col_names: str) -> list:
    """Add any missing column names to row 1 and return the full header list."""
    header_row = [
        cell.value.strip() if isinstance(cell.value, str) else cell.value
        for cell in ws[1]
    ]
    for name in col_names:
        if name not in header_row:
            ws.cell(row=1, column=len(header_row) + 1, value=name)
            header_row.append(name)
    return header_row


# ── File rename helpers (UNCHANGED) ─────────────────────────────────────────

def rename_error(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if not name.startswith("error_"):
        new_path = os.path.join(folder, "error_" + name)
        os.rename(file_path, new_path)
        return new_path
    return file_path


def rename_restore(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if name.startswith("error_"):
        new_path = os.path.join(folder, name[len("error_"):])
        os.rename(file_path, new_path)
        return new_path
    return file_path


# =============================================================================
# 6. PAYLOAD BUILDER  (UNCHANGED)
# =============================================================================

def build_header_payload(row_data: dict) -> dict:
    order_dt = to_iso_date(row_data.get("Order Date")) or to_iso_date(datetime.now())
    due_dt   = to_iso_date(row_data.get("Due Date"))   or to_iso_date(datetime.now())

    return {
        "purchaseOrderHeaders": [{
            "purchaseOrderNumber":    sv(row_data, "PO Number"),
            "domainCode":             sv(row_data, "Domain Code"),
            "supplierCode":           sv(row_data, "Supplier Code"),
            "currencyCode":           sv(row_data, "Currency"),
            "exchangeRate":           1,
            "exchangeRate2":          1,
            "exchangeRateType":       "ACCOUNTING",
            "creditTermsCode":        sv(row_data, "Credit Terms"),
            "daybookSetCode":         sv(row_data, "Daybook Set"),
            "shipToSite":             sv(row_data, "Ship To Site"),
            "billToCode":             sv(row_data, "Bill To Code"),
            "taxEnvironment":         sv(row_data, "Tax Environment", "USA"),
            "languageCode":           sv(row_data, "Language Code", "us"),
            "contact":                sv(row_data, "Contact"),
            "remarks":                sv(row_data, "Remarks"),
            "orderDate":              order_dt,
            "dueDate":                due_dt,
            "pricingDate":            order_dt,
            "lastPriceDate":          order_dt,
            "startEffective":         order_dt,
            "endEffective":           order_dt,
            "orderRevisionDate":      order_dt,
            "taxDateSummary":         order_dt,
            "orderStatus":            "O",
            "isConfirmed":            True,
            "isFixedPrice":           True,
            "isPrintPurchaseOrder":   True,
            "isTaxable":              True,
            "isDisplayTaxAmounts":    True,
            "isUsingConsignmentInventory": True,
            "isIntrastatUsed":        True,
            "isPredefaulted":         True,
            "maximumAgingDays":       90,
            "transmitFlag":           "1",
            "ERSOption":              "1",
            "ERSPriceListOption":     "0",
            "dataOperation":          "C",
            "concurrencyHash":        "",
            "disallowedActions":      "",
            "disallowedActionsMessage": "",
            "supplementaryMessages":  [],
            "POIntrastats":           [],
        }]
    }


# =============================================================================
# 7. LINE CREATION FLOW  (API sequencing UNCHANGED — only response/error
#    handling now goes through safe_get()/safe_post() + normalize_response())
# =============================================================================

def create_line(
    domain:        str,
    po_num:        str,
    line_row_data: dict,
    tm:            TokenManager,
) -> tuple[bool, list, int]:
    """
    Full multi-step line creation:
      init → fieldChange(siteCode) → fieldChange(itemCode) →
      isReceived check → fieldChange(quantityOrdered) →
      fieldChange(purchaseCost) → sync/commit

    Line number is assigned by QAD in the init response — we never override it.
    Returns (success, errors, line_number), where errors is a normalized
    list of {"fieldName", "message", "fieldValue", "code"} dicts (empty on
    success) ready for resolve_qad_errors().
    """
    site_code = sv(line_row_data, "Site Code")
    item_code = sv(line_row_data, "Item Code")
    qty       = fv(line_row_data, "Quantity Ordered")
    price     = fv(line_row_data, "Unit Price")
    due_dt    = to_iso_date(line_row_data.get("Due Date")) or to_iso_date(datetime.now())

    # 1. Initialise blank line — QAD returns the next available line number
    print("    ↳ Init line (QAD will assign number)...")
    resp, resp_json, errors = safe_get(INIT_LINE_URL.format(domain=domain, po=po_num), tm)
    if errors:
        return False, errors, 0

    lines = (resp_json or {}).get("data", {}).get("purchaseOrderLines", [])
    if not lines:
        return False, [{"fieldName": None, "message": "Init returned no line object", "fieldValue": None, "code": None}], 0

    line        = lines[0]
    line_number = line["purchaseOrderLine"]   # QAD-assigned — do not override
    line["dueDate"] = due_dt

    # 2. fieldChange: siteCode
    line["siteCode"] = site_code
    resp, resp_json, errors = safe_post(
        FIELD_CHANGE_URL.format(fieldName="siteCode"),
        {"purchaseOrderLines": [line]}, tm,
    )
    if errors:
        return False, errors, line_number
    lines = (resp_json or {}).get("data", {}).get("purchaseOrderLines", [])
    if not lines:
        return False, [{"fieldName": "siteCode", "message": "fieldChange(siteCode) returned no line", "fieldValue": site_code, "code": None}], line_number
    line = lines[0]

    # 3. fieldChange: itemCode
    line["itemCode"] = item_code
    resp, resp_json, errors = safe_post(
        FIELD_CHANGE_URL.format(fieldName="itemCode"),
        {"purchaseOrderLines": [line]}, tm,
    )
    if errors:
        return False, errors, line_number
    lines = (resp_json or {}).get("data", {}).get("purchaseOrderLines", [])
    if not lines:
        return False, [{"fieldName": "itemCode", "message": "fieldChange(itemCode) returned no line", "fieldValue": item_code, "code": None}], line_number
    line = lines[0]

    time.sleep(0.1)

    # 4. isReceived check (fire-and-forget) — use QAD-assigned line number
    safe_get(IS_RECEIVED_URL.format(domain=domain, po=po_num, line=line_number), tm)

    time.sleep(0.1)

    # 5. fieldChange: quantityOrdered
    line["quantityOrdered"] = qty
    resp, resp_json, errors = safe_post(
        FIELD_CHANGE_URL.format(fieldName="quantityOrdered"),
        {"purchaseOrderLines": [line]}, tm,
    )
    if errors:
        return False, errors, line_number
    lines = (resp_json or {}).get("data", {}).get("purchaseOrderLines", [])
    if not lines:
        return False, [{"fieldName": "quantityOrdered", "message": "fieldChange(quantityOrdered) returned no line", "fieldValue": qty, "code": None}], line_number
    line = lines[0]

    # 6. fieldChange: purchaseCost
    line["purchaseCost"] = price
    resp, resp_json, errors = safe_post(
        FIELD_CHANGE_URL.format(fieldName="purchaseCost"),
        {"purchaseOrderLines": [line]}, tm,
    )
    if errors:
        return False, errors, line_number
    lines = (resp_json or {}).get("data", {}).get("purchaseOrderLines", [])
    if not lines:
        return False, [{"fieldName": "purchaseCost", "message": "fieldChange(purchaseCost) returned no line", "fieldValue": price, "code": None}], line_number
    line = lines[0]

    # 7. Sync / commit line
    resp, resp_json, errors = safe_post(SYNC_LINE_URL, {"purchaseOrderLines": [line]}, tm)
    if errors:
        return False, errors, line_number

    return True, [], line_number


# =============================================================================
# 8. SINGLE-FILE PROCESSOR
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    print(f"\n{'─'*55}")
    print(f"  Loading: {os.path.basename(file_path)}")
    print(f"{'─'*55}")

    wb = openpyxl.load_workbook(file_path)

    if "Header" not in wb.sheetnames or "Lines" not in wb.sheetnames:
        print("ERROR: Workbook must have sheets named 'Header' and 'Lines'")
        raise RuntimeError("Workbook must have sheets named 'Header' and 'Lines'")

    ws_h = wb["Header"]
    ws_l = wb["Lines"]

    h_headers = ensure_columns(ws_h, "Status", "Error")
    l_headers = ensure_columns(ws_l, "Status", "Error")

    h_status_col = h_headers.index("Status") + 1
    h_error_col  = h_headers.index("Error")  + 1
    l_status_col = l_headers.index("Status") + 1
    l_error_col  = l_headers.index("Error")  + 1

    # ── Index header rows by PO number ────────────────────────────────────
    header_rows: dict[str, tuple[int, dict]] = {}
    for row_idx, row in enumerate(ws_h.iter_rows(min_row=2), start=2):
        row_values = [cell.value for cell in row]
        if not any(row_values):
            continue
        row_data = dict(zip(h_headers, row_values))
        po_num   = sv(row_data, "PO Number")
        if po_num:
            header_rows[po_num] = (row_idx, row_data)

    # ── Index line rows by PO number ──────────────────────────────────────
    line_rows: dict[str, list[tuple[int, dict]]] = defaultdict(list)
    for row_idx, row in enumerate(ws_l.iter_rows(min_row=2), start=2):
        row_values = [cell.value for cell in row]
        if not any(row_values):
            continue
        row_data = dict(zip(l_headers, row_values))
        po_num   = sv(row_data, "PO Number")
        if po_num:
            line_rows[po_num].append((row_idx, row_data))

    h_success = h_skip = h_fail = 0
    l_success = l_skip = l_fail = 0

    total_pos = len(header_rows)

    for idx, (po_num, (h_row_idx, h_row_data)) in enumerate(header_rows.items(), start=1):
        if _progress_callback:
            _progress_callback(idx, total_pos)

        print(f"\n  PO: {po_num}")

        # ── Step 1: Create header ─────────────────────────────────────────
        if str(h_row_data.get("Status", "")).strip().upper() == "DONE":
            print("  Header — already DONE, skipping")
            h_skip += 1
        else:
            # Mandatory check before any API call
            missing = _check_mandatory(h_row_data, MANDATORY_HEADER_COLUMNS)
            if missing:
                h_fail += 1
                mark_error(
                    ws_h, h_row_idx, h_status_col, h_error_col, h_headers,
                    missing,
                    f"Missing mandatory fields: {', '.join(missing)}",
                )
                wb.save(file_path)
                continue

            payload = build_header_payload(h_row_data)
            resp, resp_json, errors = safe_post(HEADER_URL, payload, tm)

            if errors:
                h_fail += 1
                bad_cols, error_msg = resolve_qad_errors(errors, PO_HEADER_FIELD_TO_COLUMN)
                mark_error(ws_h, h_row_idx, h_status_col, h_error_col, h_headers, bad_cols, error_msg)
                print(f"  Header FAILED: {po_num} — {error_msg}")
                wb.save(file_path)
                continue

            mark_done(ws_h, h_row_idx, h_status_col, h_error_col)
            h_success += 1
            print(f"  Header OK: {po_num}")
            wb.save(file_path)
            time.sleep(0.3)

        # ── Step 2: Create lines ──────────────────────────────────────────
        domain   = sv(h_row_data, "Domain Code")
        po_lines = line_rows.get(po_num, [])

        if not po_lines:
            print(f"  WARNING: No lines found for PO: {po_num}")
            continue

        for seq, (l_row_idx, l_row_data) in enumerate(po_lines, start=1):
            if str(l_row_data.get("Status", "")).strip().upper() == "DONE":
                print(f"    Line {seq} — already DONE, skipping")
                l_skip += 1
                continue

            missing = _check_mandatory(l_row_data, MANDATORY_LINE_COLUMNS)
            if missing:
                l_fail += 1
                mark_error(
                    ws_l, l_row_idx, l_status_col, l_error_col, l_headers,
                    missing,
                    f"Missing mandatory fields: {', '.join(missing)}",
                )
                wb.save(file_path)
                continue

            item  = sv(l_row_data, "Item Code")
            qty   = fv(l_row_data, "Quantity Ordered")
            price = fv(l_row_data, "Unit Price")
            print(f"    → {item} | qty={qty} | price={price}")

            success, errors, line_no = create_line(domain, po_num, l_row_data, tm)

            if success:
                mark_done(ws_l, l_row_idx, l_status_col, l_error_col)
                l_success += 1
                print(f"    Line {line_no} OK: {item}")
            else:
                l_fail += 1
                bad_cols, error_msg = resolve_qad_errors(errors, PO_LINE_FIELD_TO_COLUMN)
                mark_error(ws_l, l_row_idx, l_status_col, l_error_col, l_headers, bad_cols, error_msg)
                print(f"    Line {line_no} FAILED: {item} — {error_msg}")

            wb.save(file_path)
            time.sleep(0.2)

    total_success = h_success + l_success
    total_fail    = h_fail    + l_fail

    print(f"\n  Headers — OK {h_success} | Skipped {h_skip} | Failed {h_fail}")
    print(f"  Lines   — OK {l_success} | Skipped {l_skip} | Failed {l_fail}")

    return total_success, total_fail


# =============================================================================
# 9. ORCHESTRATOR
# =============================================================================

def run(folder_path: str) -> tuple[int, int]:
    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    tm = TokenManager()

    total_success = 0
    total_fail    = 0

    for file_name in xlsx_files:
        file_path = os.path.join(folder, file_name)

        s, fail = process_file(file_path, tm)
        total_success += s
        total_fail    += fail

        if fail > 0:
            file_path = rename_error(file_path)
            print(f"\n  Errors found — renamed to: {os.path.basename(file_path)}")
        elif s > 0:
            file_path = rename_restore(file_path)

        print(f"\n  Load Summary {'-'*30}")
        print(f"    Rows loaded : {s}")
        print(f"    Rows failed : {fail}")
        if fail == 0:
            print("    Result      : ALL ROWS LOADED SUCCESSFULLY ✓")
        else:
            print("    Result      : COMPLETED WITH ERRORS — fix red rows and re-run")

    print(f"\n{'='*55}")
    print(f"  TOTAL — Success: {total_success} | Failed: {total_fail}")
    print(f"{'='*55}")

    return total_success, total_fail


# =============================================================================
# 10. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["purchase_order"])
    )
    ok, fail = run(folder)
    sys.exit(0 if fail == 0 else 1)