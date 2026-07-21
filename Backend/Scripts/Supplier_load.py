"""
Supplier_load.py
----------------
Loads supplier rows from all .xlsx files in the configured supplier folder
into QAD via the supplierV2s API, then updates domain settings via mfgSuppliers.

CHANGE LOG (this revision)
---------------------------
10. MERGED: Banking is no longer a separate sheet / separate API pass.
    Bank Account Format, Supplier Bank Number, and Own Bank Number are now
    optional columns on the same Suppliers sheet. A bank entry (if present)
    is folded into the same payload/POST that creates or updates the
    supplier -- there is no second call, no "does the supplier exist yet"
    race, and no cross-sheet grouping logic. process_banking_file(),
    _read_banking_rows(), and the Banking pass in run() have been removed.
    Since a supplier can only ever have the one bank account loaded this
    way, bankNumberIsDefault and bankNumberIsActive are both forced True
    whenever a bank entry is built -- there's no Default/Active column for
    banking on the sheet.
9. Own Bank Number resolution is not a live QAD API call.
   The previous approach (resolve_bank_pay_format(), replicating the
   BankPayFormatByGLQRA browse the QAD UI uses) never actually worked --
   every call returned a near-empty, non-JSON body ("Expecting value:
   line 25 column 1" -- effectively 24 characters of nothing). Every
   Banking row with an Own Bank Number was silently failing to resolve,
   which is why bankPayFormat_ID / entityCode / ownGLCode /
   payFormatTypeCode / payFormatTypePaymentInstrument never landed.
   In this environment the Own Bank Number and its Payment Format
   are constant -- there is no need for a dynamic lookup at all. Replaced
   with a static reference table, KNOWN_OWN_BANK_NUMBERS, confirmed
   against real QAD records fetched via GET supplierV2s after a human
   saved them correctly through the UI. If a row's Own Bank Number isn't
   in the table, the row is marked ERROR telling you to add it -- it is
   never guessed.
8. FIXED (found while rebuilding #9): the bank entry payload was writing
   the resolved currency to a field called "ownBankCurrency". The real
   QAD field, confirmed against an actual saved record, is
   "ownBankCurrencyCode". Even on a lookup that resolved correctly, that
   field would have gone nowhere on save. Corrected in _build_bank_entry.

CHANGE LOG (carried over from previous revisions)
---------------------------------------------------
1-7. Full request/response debug logging to console + a timestamped log
   file per workbook; row-level try/except for full tracebacks; forced
   default bank account when a supplier has none yet (QAD error 341).

Behaviour
------------------------------------
- Fetches one OAuth token per run; refreshes only on 401
- Skips rows with Status = DONE (for Create operations only)
- Lightweight mandatory-field check before any API call
- Success check: submitResult.success == True
- Error messages extracted from submitResult.errors[].message
- On success  -> clear all red fills, clear Error, set Status = DONE
- On failure  -> red-fill first col + bad cols + Error col, log message, set Status = ERROR
- Processes every row regardless of individual failures
- Returns (ok_count, fail_count)

CREATE flow (Data Operation = C):
  1. POST to supplierV2s   -> creates supplier (bank entry, if present on
     the row, is built into the same payload's bankNumberRefV2s)
  2. GET  mfgSuppliers     -> fetch auto-created domain record
  3. PATCH + POST          -> update siteCode + daybookSetCode
  If step 3 fails -> row marked ERROR with "Supplier created but domain settings failed: ..."

UPDATE flow (Data Operation = U):
  1. GET  supplierV2s      -> fetch existing record
  2. Patch editable fields, including the bank entry (update in place if
     one already exists on the record, add fresh if not)
  3. POST supplierV2s      -> update supplier
  (Domain settings NOT touched on UPDATE -- one-time setup only)

BANK FIELDS on the Suppliers sheet (optional, no separate sheet/pass):
  - Bank Account Format, Supplier Bank Number: both blank -> no bank
    account for this supplier, skipped silently. One filled without the
    other -> treated as a data-entry mistake, row errors out.
  - Own Bank Number (optional): if populated, looked up in
    KNOWN_OWN_BANK_NUMBERS (static table) to get the GL Account /
    Payment Format / Payment Instrument bundle. Not in the table -> row
    errors out asking for the bundle to be added; never guessed.
  - bankNumberIsDefault / bankNumberIsActive are always forced True when
    a bank entry is present -- no sheet columns for either.
"""

import os
import sys
import time
import json
import datetime
import traceback
import requests
import openpyxl
from openpyxl.styles import PatternFill

# -- Path / config -----------------------------------------------------------
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT_DIR)
from config import CONFIG

# -- Fill constants -----------------------------------------------------------
RED_FILL    = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
CLEAR_FILL  = PatternFill(fill_type=None)

# -- Mandatory columns (supplier) ---------------------------------------------
MANDATORY_COLUMNS = [
    "Supplier",
    "Shared Set",
    "Business Relation",
    "Active",
    "Currency",
    "Credit Terms",
    "Invoice Status",
    "Invoice Control GL Profile",
    "Credit Note Control GL Profile",
    "Prepayment Control GL Profile",
    "Purchase Account GL Profile",
]

# -- Mandatory columns for domain settings (CREATE only) -----------------------
MANDATORY_DOMAIN_COLUMNS = [
    "Domain",
    "Site Code",
    "Daybook Set",
]

# -- Bank columns (all optional; presence/absence is checked in code, not
#    listed in MANDATORY_COLUMNS above) -----------------------------------
BANK_ACCOUNT_FORMAT_COL = "Bank Account Format"
BANK_ACCOUNT_NUMBER_COL = "Supplier Bank Number"
OWN_BANK_NUMBER_COL     = "Own Bank Number"

# -- Bank number formatting rules, per Bank Account Format code ---------------
# Confirmed against two captured examples (AU format, 6+rest split, space
# delimiter). Add more entries here as other formats are confirmed.
BANK_FORMAT_SECTIONS = {
    "AU": [6, None],   # 6-digit section, then remaining digits. None = "rest".
}
BANK_FORMAT_DELIMITER = " "


def _split_bank_number(bank_format: str, raw_digits: str) -> list[str]:
    """Split raw bank digits into sections per BANK_FORMAT_SECTIONS.
    Falls back to a single unsplit section for unknown formats."""
    layout = BANK_FORMAT_SECTIONS.get(bank_format)
    if not layout:
        return [raw_digits]

    parts, pos = [], 0
    for length in layout:
        if length is None:
            parts.append(raw_digits[pos:])
            pos = len(raw_digits)
        else:
            parts.append(raw_digits[pos:pos + length])
            pos += length
    return parts


# =============================================================================
# 0. DEBUG LOGGING
# =============================================================================

class DebugLog:
    def __init__(self, folder: str, sheet_label: str):
        ts = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
        self.path = os.path.join(folder, f"{sheet_label}_{ts}.log")
        self._fh = open(self.path, "a", encoding="utf-8")
        self.write(f"=== Debug log started {ts} ===")

    def write(self, msg: str):
        line = f"[{datetime.datetime.now().strftime('%H:%M:%S')}] {msg}"
        print(line)
        self._fh.write(line + "\n")
        self._fh.flush()

    def request(self, method: str, url: str, params=None, payload=None):
        self.write(f"REQUEST {method} {url}")
        if params:
            self.write(f"  params: {json.dumps(params, default=str)}")
        if payload is not None:
            self.write(f"  payload: {json.dumps(payload, default=str)[:4000]}")

    def response(self, status: int, body):
        self.write(f"RESPONSE HTTP {status}")
        try:
            self.write(f"  body: {json.dumps(body, default=str)[:4000]}")
        except Exception:
            self.write(f"  body (unserializable): {str(body)[:4000]}")

    def close(self):
        self.write("=== Debug log closed ===")
        self._fh.close()


_LOG: DebugLog | None = None


def _log(msg: str):
    if _LOG is not None:
        _LOG.write(msg)
    else:
        print(msg)


# =============================================================================
# 1. AUTHENTICATION
# =============================================================================

def _fetch_token() -> str:
    url = f"{CONFIG['qad']['base_url']}/oauth/token"
    _log(f"AUTH requesting new token from {url}")
    resp = requests.post(url, data=CONFIG["qad"]["auth"], timeout=30)
    resp.raise_for_status()
    token = resp.json().get("access_token")
    if not token:
        _log("AUTH FAILED -- response had no access_token")
        raise RuntimeError("OAuth response did not contain access_token")
    _log("AUTH token acquired")
    return token


class TokenManager:
    """Holds one token for the run; refreshes on demand (401)."""

    def __init__(self):
        self._token: str | None = None

    def get(self) -> str:
        if self._token is None:
            self._token = _fetch_token()
        return self._token

    def refresh(self) -> str:
        _log("AUTH token refresh triggered (401 received)")
        self._token = _fetch_token()
        return self._token


class _TokenExpired(Exception):
    pass


# =============================================================================
# 2. OWN BANK NUMBER REFERENCE DATA  (static -- no live QAD lookup)
# =============================================================================
# The GL Account / Payment Format / Payment Instrument bundle behind an Own
# Bank Number is NOT independently settable through the supplierV2s API --
# QAD's UI resolves it from a picker at data-entry time and stamps the
# resolved values onto the bank entry before it's ever saved. There is no
# server-side derivation on save: whatever bundle you send is what's stored.
#
# To add a new Own Bank Number: save one bank entry through the QAD UI
# using that Own Bank Number, GET that supplier's supplierV2s record, and
# copy bankPayFormat_ID / entityCode / ownBankCurrencyCode / ownGLCode /
# payFormatTypeCode / payFormatTypePaymentInstrument from its
# bankNumberRefV2s entry into this table. Do not guess these values.

KNOWN_OWN_BANK_NUMBERS: dict[str, dict] = {
    "701000009527": {
        "bankPayFormat_ID":               1219554787,
        "entityCode":                     "PIIEM01",
        "ownBankCurrencyCode":            "INR",
        "ownGLCode":                      "11101426",
        "payFormatTypeCode":              "AP-BNK",
        "payFormatTypePaymentInstrument": "ELECTRONIC",
    },
}


def resolve_own_bank_number(own_bank_number: str) -> dict | None:
    """Static lookup -- returns the confirmed bundle for a known Own Bank
    Number, or None if it hasn't been added to KNOWN_OWN_BANK_NUMBERS yet."""
    bundle = KNOWN_OWN_BANK_NUMBERS.get(own_bank_number)
    if bundle is None:
        _log(f"LOOKUP no entry in KNOWN_OWN_BANK_NUMBERS for '{own_bank_number}'")
    else:
        _log(f"LOOKUP resolved '{own_bank_number}' from static table -> {bundle}")
    return bundle


# =============================================================================
# 3. BANK ENTRY BUILDER
# =============================================================================

def _cell_str(row_data: dict, col: str) -> str:
    """Safely read a cell value as a trimmed string. openpyxl returns Python
    None for an empty cell -- row_data.get(col, "") does NOT fall back to
    "" in that case (the key exists, its value is just None), so a naive
    str(row_data.get(col, "")) turns an empty cell into the literal text
    "None". Every read of a sheet cell outside build_payload's own val()
    helper must go through this instead."""
    v = row_data.get(col)
    return str(v).strip() if v is not None else ""


def _bank_fields_present(row_data: dict) -> tuple[bool, list[str]]:
    """
    Determine whether this row has a bank account to build, and catch the
    half-filled case (one of the two required bank fields present without
    the other) as a data-entry mistake rather than silently skipping it.

    Returns (has_bank_entry, missing_fields_if_partial).
      - Both blank            -> (False, [])           -- no bank account, skip silently
      - Both filled           -> (True,  [])           -- build the entry
      - Exactly one filled    -> (False, [<missing>])  -- error, not a skip
    """
    fmt    = _cell_str(row_data, BANK_ACCOUNT_FORMAT_COL)
    number = _cell_str(row_data, BANK_ACCOUNT_NUMBER_COL)

    if not fmt and not number:
        return False, []

    if fmt and number:
        return True, []

    missing = []
    if not fmt:
        missing.append(BANK_ACCOUNT_FORMAT_COL)
    if not number:
        missing.append(BANK_ACCOUNT_NUMBER_COL)
    return False, missing


def _build_bank_entry(row_data: dict, shared_set: str, supplier_code: str, sequence: int) -> tuple[dict | None, str]:
    """
    Build one bankNumberRefV2s entry from the Suppliers-sheet row's bank
    columns. Returns (entry, error_msg) -- entry is None if Own Bank Number
    was supplied but not found in KNOWN_OWN_BANK_NUMBERS (error_msg set in
    that case). bankNumberIsDefault and bankNumberIsActive are always
    forced True -- a supplier loaded this way only ever has one bank
    account, so "default" and "active" aren't optional choices.
    """
    bank_format = _cell_str(row_data, BANK_ACCOUNT_FORMAT_COL)
    raw_digits  = _cell_str(row_data, BANK_ACCOUNT_NUMBER_COL).replace(" ", "")
    parts       = _split_bank_number(bank_format, raw_digits)
    formatted   = BANK_FORMAT_DELIMITER.join(parts)

    own_bank_number = _cell_str(row_data, OWN_BANK_NUMBER_COL)
    bundle = None
    if own_bank_number:
        bundle = resolve_own_bank_number(own_bank_number)
        if bundle is None:
            return None, (
                f"Own Bank Number '{own_bank_number}' is not in "
                f"KNOWN_OWN_BANK_NUMBERS. Add its GL Account / Payment "
                f"Format / Payment Instrument bundle to Supplier_load.py "
                f"(pull it from a supplier already saved correctly in QAD) "
                f"before this row can be loaded."
            )

    section_refs = [
        {
            "sharedSetCode":             shared_set,
            "supplierCode":              supplier_code,
            "bankNumberSequence":        sequence,
            "bankNumberSectionSequence": i + 1,
            "sectionValue":              part,
        }
        for i, part in enumerate(parts)
    ]

    entry = {
        "bankAccFormatCode":              bank_format,
        "bankBusinessRelationCode":       "",
        "bankNumber":                     raw_digits,
        "bankNumberFormatted":            formatted,
        "bankNumberBankAccountType":      "",
        "bankNumberBranch":               "",
        "bankNumberSwiftCode":            "",
        "bankNumberIsActive":             True,
        "bankNumberIsDefault":            True,
        "bankNumberSequence":             sequence,
        "sharedSetCode":                  shared_set,
        "supplierCode":                   supplier_code,
        "bankNumberPayCodeRefV2s":        [],
        "bankNumberSectionRefV2s":        section_refs,
        # -- Own Bank Number linkage bundle -- all blank unless resolved ----
        # NOTE: field is "ownBankCurrencyCode", not "ownBankCurrency" --
        # confirmed against a real saved QAD record.
        "bankPayFormat_ID":               0,
        "entityCode":                     "",
        "ownBankNumber":                  "",
        "ownBankCurrencyCode":            "",
        "ownGLCode":                      "",
        "payFormatTypeCode":              "",
        "payFormatTypePaymentInstrument": "",
    }

    if bundle:
        entry["bankPayFormat_ID"]               = bundle.get("bankPayFormat_ID", 0)
        entry["entityCode"]                     = bundle.get("entityCode", "")
        entry["ownBankNumber"]                  = own_bank_number
        entry["ownBankCurrencyCode"]            = bundle.get("ownBankCurrencyCode", "")
        entry["ownGLCode"]                      = bundle.get("ownGLCode", "")
        entry["payFormatTypeCode"]              = bundle.get("payFormatTypeCode", "")
        entry["payFormatTypePaymentInstrument"] = bundle.get("payFormatTypePaymentInstrument", "")

    return entry, ""


def _apply_bank_entry_to_supplier(supplier: dict, row_data: dict) -> tuple[bool, str]:
    """
    Given a supplier dict (already containing whatever bankNumberRefV2s it
    currently has -- [] for a brand-new Create, or the fetched list for an
    Update), add or patch the single bank entry described by this row's
    bank columns in place. Returns (ok, error_msg).

    - No bank fields on the row -> no-op, ok=True.
    - Partial bank fields        -> ok=False, error_msg set.
    - Full bank fields, no existing entry -> append one new entry.
    - Full bank fields, existing entry present -> patch that entry's
      fields in place (uri/tcRowid/bankNumber_ID preserved) so QAD treats
      it as a modification, not a duplicate.
    """
    has_bank, missing = _bank_fields_present(row_data)

    if missing:
        return False, f"Incomplete bank account: missing {', '.join(missing)}"

    if not has_bank:
        return True, ""

    shared_set    = supplier.get("sharedSetCode", "")
    supplier_code = supplier.get("supplierCode", "")
    existing_list = supplier.get("bankNumberRefV2s") or []

    if existing_list:
        # Exactly one bank account ever loaded this way -- patch the first
        # (only) existing entry in place, keeping its identity fields.
        existing_entry = existing_list[0]
        new_entry, err = _build_bank_entry(row_data, shared_set, supplier_code, existing_entry.get("bankNumberSequence", 1))
        if new_entry is None:
            return False, err

        # Preserve identity fields so QAD treats this as an update to the
        # same bank record, not a new one.
        for identity_field in ("uri", "tcRowid", "bankNumber_ID"):
            if identity_field in existing_entry:
                new_entry[identity_field] = existing_entry[identity_field]

        existing_list[0] = new_entry
        supplier["bankNumberRefV2s"] = existing_list
    else:
        new_entry, err = _build_bank_entry(row_data, shared_set, supplier_code, 1)
        if new_entry is None:
            return False, err
        supplier["bankNumberRefV2s"] = [new_entry]

    return True, ""


def _verify_bank_entry(shared_set: str, supplier_code: str, expected_entry: dict, token: str) -> list[str]:
    """
    Post-write verification for the single bank entry just sent. Re-GETs
    the supplier and checks the entry with matching bankNumberFormatted
    actually exists, and (if we sent a non-zero bankPayFormat_ID) that it
    persisted rather than silently landing at 0.
    """
    problems = []

    verify_resp = get_supplier(shared_set, supplier_code, token)
    supplier_list = verify_resp.get("data", {}).get("supplierV2s")
    if not supplier_list:
        return [f"Verification GET failed for {shared_set}.{supplier_code} -- could not confirm write"]

    stored_entries = supplier_list[0].get("bankNumberRefV2s") or []
    stored = next(
        (e for e in stored_entries if e.get("bankNumberFormatted") == expected_entry["bankNumberFormatted"]),
        None,
    )

    if stored is None:
        problems.append(f"Bank entry '{expected_entry['bankNumberFormatted']}' not found in QAD after save")
        return problems

    if expected_entry["bankPayFormat_ID"] and not stored.get("bankPayFormat_ID"):
        problems.append(
            f"Bank entry '{expected_entry['bankNumberFormatted']}': Own Bank Number linkage did not "
            f"persist (sent bankPayFormat_ID={expected_entry['bankPayFormat_ID']}, "
            f"QAD stored bankPayFormat_ID=0 / ownBankNumber blank)"
        )

    if expected_entry.get("ownBankCurrencyCode") and not stored.get("ownBankCurrencyCode"):
        problems.append(
            f"Bank entry '{expected_entry['bankNumberFormatted']}': ownBankCurrencyCode did not persist "
            f"(sent '{expected_entry['ownBankCurrencyCode']}', QAD stored blank)"
        )

    return problems


# =============================================================================
# 4. PAYLOAD BUILDER  (explicit -- every field visible)
# =============================================================================

def build_payload(row: dict) -> dict:
    def val(col: str) -> str:
        v = row.get(col)
        return str(v).strip() if v is not None else ""

    def bool_val(col: str) -> bool:
        return val(col).lower() == "yes"

    def float_val(col: str) -> float:
        return float(val(col) or 0)

    def int_val(col: str) -> int:
        return int(float(val(col) or 0))

    supplier_code = val("Supplier")
    shared_set    = val("Shared Set")
    uri           = f"urn:be:com.qad.base.supplier.ISupplierV2:{shared_set}.{supplier_code}"

    return {
        "supplementaryMessages": [],
        "supplierV2s": [
            {
                # -- URIs / Identity --------------------------------------------
                "uri":                                      uri,
                "instanceURI":                              uri,
                "supplierCode":                             supplier_code,
                "supplierID":                               0,
                "sharedSetCode":                            shared_set,
                "sharedSetID":                              0,

                # -- Business Relation -------------------------------------------
                "businessRelationCode":                     val("Business Relation"),
                "businessRelationID":                       0,
                "businessRelationName":                     val("Name"),
                "businessRelationName2":                    val("Second Name"),
                "businessRelationName3":                    val("Third Name"),
                "businessRelationConcurrencyHash":          "",
                "isBusinessRelationActive":                 True,
                "isBusinessRelationFieldsEnabled":          True,
                "isBusinessRelationIntercompany":           False,
                "isCreateBusinessRelationRequired":         True,
                "isInternalEntity":                         False,
                "intercompanyCode":                         "",

                # -- Address ------------------------------------------------------
                "addressID":                                0,
                "addressName":                              val("Name"),
                "addressSearchName":                        val("Search Name"),
                "addressTypeCode":                          val("Address Type"),
                "addressConcurrencyHash":                   "",
                "street1":                                  val("Address 1"),
                "street2":                                  val("Address 2"),
                "street3":                                  "",
                "city":                                     val("City"),
                "cityCode":                                 "",
                "zipCode":                                  val("Postal Code"),
                "stateCode":                                val("State"),
                "stateDescription":                         "",
                "stateTax":                                 val("State Tax"),
                "countryCode":                              val("Country"),
                "countryDescription":                       "",
                "countyCode":                               "",
                "countyDescription":                        "",
                "postalFormat":                             "0",
                "latitude":                                 0,
                "longitude":                                0,
                "isTemporaryAddress":                       False,

                # -- Status / Concurrency ------------------------------------------
                "changeStatus":                             "2",
                "dataOperation":                            "",
                "concurrencyHash":                          "",
                "disallowedActions":                        "",
                "disallowedActionsMessage":                 "",
                "isPredefaulted":                           False,
                "isDomainRestricted":                       False,
                "lastModifiedDate":                         "",
                "lastModifiedTime":                         0,
                "lastModifiedUser":                         "",

                # -- Active / Type --------------------------------------------------
                "isActive":                                 bool_val("Active"),
                "supplierTypeCode":                         val("Supplier Type"),
                "supplierTypeID":                           0,
                "supplierTypeDescription":                  "",
                "purchaseTypeCode":                         val("Purchase Type"),
                "purchaseTypeID":                           0,
                "purchaseTypeDescription":                  "",
                "purchaseCodeID":                           0,

                # -- Currency / Language -----------------------------------------
                "currencyCode":                             val("Currency"),
                "currencyDescription":                      "",
                "currencyID":                               0,
                "languageCode":                             val("Language"),
                "languageDescription":                      "",

                # -- Credit Terms -------------------------------------------------
                "creditTermsCode":                          val("Credit Terms"),
                "creditTermsID":                            0,
                "creditTermsDescription":                   "",
                "normalCreditTermsID":                      0,
                "normalCreditTermsType":                    "",

                # -- Invoice / Status ----------------------------------------------
                "invoiceStatusCode":                        val("Invoice Status"),
                "invoiceStatusDescription":                 "",
                "invoiceStatusID":                          0,
                "isIndividualPayments":                     False,
                "isSplitAccount":                           False,

                # -- GL Profiles ---------------------------------------------------
                "invoiceControlGLProfileCode":              val("Invoice Control GL Profile"),
                "invoiceControlGLProfileDesc":              "",
                "invoiceControlGLProfileID":                0,
                "creditNoteControlGLProfileCode":           val("Credit Note Control GL Profile"),
                "creditNoteControlGLProfileDesc":           "",
                "creditNoteControlGLProfileID":             0,
                "prePaymentControlGLProfileCode":           val("Prepayment Control GL Profile"),
                "prePaymentControlGLProfileDesc":           "",
                "prePaymentControlGLProfileID":             0,
                "purchaseAccountGLProfileCode":             val("Purchase Account GL Profile"),
                "purchaseAccountProfileDesc":               "",
                "purchaseAccountGLProfileID":               0,
                "financeChargeGLProfileID":                 0,

                # -- Tax -------------------------------------------------------------
                "taxZone":                                  val("Tax Zone"),
                "taxZoneDescription":                       "",
                "taxClass":                                 val("Tax Class" or ""),
                "taxClassDescription":                      "",
                "taxDeclaration":                           0,
                "taxUsage":                                 val("Tax Usuage"),   # note: matches Excel typo
                "taxUsageDescription":                      "",
                "taxLevel":                                 "",
                "taxNature":                                "SERVICE",
                "taxIDFiscalCode":                          "",
                "foreignFiscalCode":                        "",
                "chamberOfCommerceNumber":                  "",
                "birthCity":                                "",
                "TIDNotice":                                "",
                "whtCertFormatCode":                        "",
                "whtCertFormatDescription":                 "",
                "whtCertFormatID":                          0,
                "isTaxableSupplier":                        bool_val("Taxable (Yes/No)"),
                "isTaxInCity":                              bool_val("Tax in City(Yes/No)"),
                "isTaxIncluded":                             bool_val("Tax Included(Yes/No)"),
                "isTaxReport":                              bool_val("Tax Report"),
                "isTaxReportForBusinessRelation":           False,
                "isTaxConfirmed":                           False,
                "isWithholdingTax":                         False,
                "isLastFiling":                             False,
                "isReportedIN":                             False,
                "federalTax":                               val("Federal tax"),
                "stateTax":                                 val("State Tax"),
                "miscellaneousTax1":                        val("Miscellaneous Tax 1"),
                "miscellaneousTax2":                        val("Miscellaneous Tax 2"),
                "miscellaneousTax3":                        val("Miscellaneous Tax 3"),
                "nameControl":                              "",
                "EORINumber":                               "",

                # -- Payment ------------------------------------------------------
                "paymentGroupCode":                         "",
                "paymentGroupDescription":                  "",
                "paymentGroupID":                           0,
                "isPayBankCharge":                          False,
                "isCompensationAllowed":                    False,
                "externalCustomerNumber":                   "",
                "deliveryConditionID":                      0,

                # -- Remittance -----------------------------------------------------
                "isRemittanceRequired":                     False,
                "isSendRemittance":                         False,
                "remittanceAddressChangeStatus":            "",
                "remittanceAddressTypeCode":                "",
                "remitSupplierContactV2s":                  [],

                # -- Corporate Group ------------------------------------------------
                "corporateGroupCode":                       "",
                "corporateGroupDescription":                "",
                "corporateGroupID":                         0,

                # -- Sub Account ----------------------------------------------------
                "subAccountProfileCode":                    "",
                "subAccountProfileDesc":                    "",
                "subAccountProfileID":                      0,

                # -- Contact ----------------------------------------------------------
                "EMail":                                    val("Email"),
                "fax":                                      "",
                "telephone":                                val("Telephone"),
                "webSite":                                  "",
                "commentNote":                               "",
                "creditAgencyReference":                    "",

                # -- Custom Fields --------------------------------------------------
                "customShort0":  "", "customShort1":  "", "customShort2":  "",
                "customShort3":  "", "customShort4":  "", "customShort5":  "",
                "customShort6":  "", "customShort7":  "", "customShort8":  "",
                "customShort9":  "", "customShort10": "", "customShort11": "",
                "customShort12": "", "customShort13": "", "customShort14": "",
                "customShort15": "", "customShort16": "", "customShort17": "",
                "customShort18": "", "customShort19": "",
                "customLong0":   "", "customLong1":   "",
                "customNote":    "",
                "customCombo0":  "", "customCombo1":  "", "customCombo2":  "",
                "customCombo3":  "", "customCombo4":  "", "customCombo5":  "",
                "customCombo6":  "", "customCombo7":  "", "customCombo8":  "",
                "customCombo9":  "", "customCombo10": "", "customCombo11": "",
                "customCombo12": "B2B", "customCombo13": "", "customCombo14": "",
                "customDecimal0": 0, "customDecimal1": 0, "customDecimal2": 0,
                "customDecimal3": 0, "customDecimal4": 0,
                "customInteger0": 0, "customInteger1": 0, "customInteger2": 0,
                "customInteger3": 0, "customInteger4": 0,

                # -- Sub-lists ----------------------------------------------------
                "supplierContactV2s":                       [],
                "supplierSafDefaultV2s":                    [],
                "supplierVatV2s":                           [],
                "bankNumberRefV2s":                         [],
                "MDMBankNrSharedSets":                      [],
            }
        ],
    }


# =============================================================================
# 5. SUPPLIER API
# =============================================================================

def post_supplier(payload: dict, token: str, is_create: bool = False) -> tuple[bool, str]:
    """
    POST supplier payload to QAD.
    Returns (success, error_msg). Raises _TokenExpired on 401.

    CREATE -> viewUri only  (no sharedSetCode/supplierCode in query string)
    UPDATE -> sharedSetCode + supplierCode + viewUri
    """
    supplier      = payload["supplierV2s"][0]
    shared_set    = supplier.get("sharedSetCode", "")
    supplier_code = supplier.get("supplierCode", "")

    if is_create:
        url = (
            f"{CONFIG['qad']['base_url']}/api/erp/supplierV2s"
            f"?viewUri=urn:be:com.qad.base.supplier.ISupplierV2"
        )
    else:
        url = (
            f"{CONFIG['qad']['base_url']}/api/erp/supplierV2s"
            f"?sharedSetCode={shared_set}&supplierCode={supplier_code}"
            f"&viewUri=urn:be:com.qad.base.supplier.ISupplierV2"
        )

    _log(f"POST supplier {shared_set}.{supplier_code} (create={is_create})")
    if _LOG is not None:
        _LOG.request("POST", url, payload=payload)

    resp = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type":  "application/json",
        },
        json=payload,
        timeout=30,
    )

    if resp.status_code == 401:
        _log(f"POST supplier {shared_set}.{supplier_code} -> HTTP 401, will refresh token")
        raise _TokenExpired()

    resp_json = resp.json()
    if _LOG is not None:
        _LOG.response(resp.status_code, resp_json)

    submit = resp_json.get("submitResult", {})

    if submit.get("success") is True:
        _log(f"POST supplier {shared_set}.{supplier_code} -> success")
        return True, ""

    errors = submit.get("errors", [])

    _log(f"POST supplier {shared_set}.{supplier_code} FAILED -- HTTP {resp.status_code}")
    _log(json.dumps(submit, indent=2, default=str))

    error_msg = "; ".join(
        e.get("message", "").strip()
        for e in errors
        if e.get("message", "").strip()
    ) or f"HTTP {resp.status_code} — submitResult.success was not True"

    return False, error_msg


def get_supplier(shared_set: str, supplier_code: str, token: str) -> dict:
    url = f"{CONFIG['qad']['base_url']}/api/erp/supplierV2s"
    params = {
        "sharedSetCode": shared_set,
        "supplierCode":  supplier_code,
        "viewUri":       "urn:be:com.qad.base.supplier.ISupplierV2",
    }

    if _LOG is not None:
        _LOG.request("GET", url, params=params)

    resp = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}"},
        params=params,
        timeout=30,
    )

    if resp.status_code == 401:
        _log(f"GET supplier {shared_set}.{supplier_code} -> HTTP 401, will refresh token")
        raise _TokenExpired()

    resp_json = resp.json()
    if _LOG is not None:
        _LOG.response(resp.status_code, resp_json)

    if not resp_json.get("data") or not resp_json["data"].get("supplierV2s"):
        _log(f"GET supplier {shared_set}.{supplier_code} FAILED -- HTTP {resp.status_code}, empty/missing data")

    return resp_json


# =============================================================================
# 6. DOMAIN SETTINGS API  (mfgSuppliers)
# =============================================================================

def get_mfg_supplier(domain: str, supplier_code: str, token: str) -> dict:
    """Fetch the auto-created mfgSupplier domain record."""
    url = f"{CONFIG['qad']['base_url']}/api/erp/mfgSuppliers"
    params = {
        "domainContext": domain,
        "supplierCode":  supplier_code,
        "viewUri":       "urn:be:com.qad.base.supplier.IMfgSupplier",
    }

    if _LOG is not None:
        _LOG.request("GET", url, params=params)

    resp = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}"},
        params=params,
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    resp_json = resp.json()
    if _LOG is not None:
        _LOG.response(resp.status_code, resp_json)

    mfg_list = resp_json.get("data", {}).get("mfgSuppliers") or resp_json.get("mfgSuppliers")
    if not mfg_list:
        _log(f"GET mfg_supplier domain={domain} supplier={supplier_code} FAILED -- HTTP {resp.status_code}")

    return resp_json


def post_mfg_supplier(payload: dict, domain: str, supplier_code: str, token: str) -> tuple[bool, str]:
    """POST updated domain settings back to QAD."""
    url = (
        f"{CONFIG['qad']['base_url']}/api/erp/mfgSuppliers"
        f"?domainContext={domain}&supplierCode={supplier_code}"
        f"&viewUri=urn:be:com.qad.base.supplier.IMfgSupplier"
    )

    if _LOG is not None:
        _LOG.request("POST", url, payload=payload)

    resp = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type":  "application/json",
        },
        json=payload,
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    resp_json = resp.json()
    if _LOG is not None:
        _LOG.response(resp.status_code, resp_json)

    submit = resp_json.get("submitResult", {})

    if submit.get("success") is True:
        return True, ""

    errors = submit.get("errors", [])

    _log(f"POST mfg_supplier domain={domain} supplier={supplier_code} FAILED -- HTTP {resp.status_code}")
    _log(json.dumps(submit, indent=2, default=str))

    error_msg = "; ".join(
        e.get("message", "").strip()
        for e in errors
        if e.get("message", "").strip()
    ) or f"HTTP {resp.status_code} — submitResult.success was not True"

    return False, error_msg


def update_domain_settings(
    supplier_code: str,
    domain:        str,
    site_code:     str,
    daybook_set:   str,
    token:         str,
) -> tuple[bool, str]:
    """
    GET auto-created mfgSupplier -> patch siteCode + daybookSetCode -> POST back.
    Returns (success, error_msg).
    """
    existing = get_mfg_supplier(domain, supplier_code, token)

    mfg_list = (
        existing.get("data", {}).get("mfgSuppliers")
        or existing.get("mfgSuppliers")
    )

    if not mfg_list:
        return False, f"Domain settings record not found for supplier '{supplier_code}' in domain '{domain}'"

    payload    = existing.get("data", existing)
    mfg_record = mfg_list[0]

    mfg_record["siteCode"]       = site_code
    mfg_record["daybookSetCode"] = daybook_set

    return post_mfg_supplier(payload, domain, supplier_code, token)


# =============================================================================
# 7. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict, columns: list) -> list[str]:
    missing = []
    for col in columns:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 8. WORKBOOK HELPERS
# =============================================================================

def _mark_done(ws, row_idx: int, status_col: int, error_col: int, note: str = ""):
    for cell in ws[row_idx]:
        cell.fill = CLEAR_FILL
    ws.cell(row=row_idx, column=status_col, value="DONE")
    ws.cell(row=row_idx, column=error_col,  value=note)


def _mark_verify(ws, row_idx: int, status_col: int, error_col: int, note: str):
    """Saved successfully (submitResult.success == True) but post-write
    verification found a mismatch -- flagged yellow so it's not mistaken
    for a clean DONE, but not red/ERROR either since the write itself did
    not fail."""
    for cell in ws[row_idx]:
        cell.fill = CLEAR_FILL
    ws.cell(row=row_idx, column=status_col).fill = YELLOW_FILL
    ws.cell(row=row_idx, column=error_col).fill  = YELLOW_FILL
    ws.cell(row=row_idx, column=status_col, value="DONE (VERIFY)")
    ws.cell(row=row_idx, column=error_col,  value=note)


def _mark_error(
    ws,
    row_idx:    int,
    status_col: int,
    error_col:  int,
    header_row: list,
    bad_cols:   list,
    error_msg:  str,
):
    for cell in ws[row_idx]:
        cell.fill = CLEAR_FILL

    ws.cell(row=row_idx, column=1).fill          = RED_FILL
    ws.cell(row=row_idx, column=status_col).fill = RED_FILL
    ws.cell(row=row_idx, column=error_col).fill  = RED_FILL
    ws.cell(row=row_idx, column=status_col, value="ERROR")
    ws.cell(row=row_idx, column=error_col,  value=error_msg)

    for col_name in bad_cols:
        if col_name in header_row:
            ws.cell(row=row_idx, column=header_row.index(col_name) + 1).fill = RED_FILL


# =============================================================================
# 9. SUPPLIER SHEET PROCESSOR  (single pass -- supplier + bank fields together)
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    """Open one workbook, process every row (supplier + optional bank
    entry together, one POST per row), return (ok_count, fail_count)."""

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # Row 1 = section labels, Row 2 = column headers, Row 3+ = data
    header_row = [
        str(c.value).strip() if c.value is not None else ""
        for c in ws[2]
    ]

    for col_name in ("Status", "Error"):
        if col_name not in header_row:
            ws.cell(row=2, column=len(header_row) + 1, value=col_name)
            header_row.append(col_name)

    status_col = header_row.index("Status") + 1
    error_col  = header_row.index("Error")  + 1

    ok_count   = 0
    fail_count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=3), start=3):

        row_values = [cell.value for cell in row]

        if not any(row_values):
            continue

        row_data = dict(zip(header_row, row_values))

        try:
            status         = _cell_str(row_data, "Status").upper()
            data_operation = _cell_str(row_data, "Data Operation").upper() or "C"

            # Skip completed CREATE rows only
            if status == "DONE" and data_operation == "C":
                continue

            _log(f"--- row {row_idx}: Supplier='{row_data.get('Supplier','')}' op={data_operation} ---")

            # -- Mandatory check --------------------------------------------------
            missing = _check_mandatory(row_data, MANDATORY_COLUMNS)
            if missing:
                fail_count += 1
                msg = f"Missing mandatory fields: {', '.join(missing)}"
                _log(f"row {row_idx} SKIPPED -- {msg}")
                _mark_error(ws, row_idx, status_col, error_col, header_row, missing, msg)
                wb.save(file_path)
                continue

            # -- Operation validity check ------------------------------------------
            if data_operation not in {"C", "U"}:
                fail_count += 1
                _mark_error(
                    ws, row_idx, status_col, error_col, header_row,
                    ["Data Operation"],
                    "Data Operation must be C, U, or blank",
                )
                wb.save(file_path)
                continue

            # -- Bank fields pre-check (partial-entry guard) ------------------------
            _, bank_missing = _bank_fields_present(row_data)
            if bank_missing:
                fail_count += 1
                msg = f"Incomplete bank account: missing {', '.join(bank_missing)}"
                _log(f"row {row_idx} SKIPPED -- {msg}")
                _mark_error(ws, row_idx, status_col, error_col, header_row, bank_missing, msg)
                wb.save(file_path)
                continue

            # -- Domain settings mandatory check (CREATE only) ---------------------
            is_create = (data_operation == "C")

            if is_create:
                missing_domain = _check_mandatory(row_data, MANDATORY_DOMAIN_COLUMNS)
                if missing_domain:
                    fail_count += 1
                    _mark_error(
                        ws, row_idx, status_col, error_col, header_row,
                        missing_domain,
                        f"Missing domain setting fields: {', '.join(missing_domain)}",
                    )
                    wb.save(file_path)
                    continue

            # -- Build payload --------------------------------------------------------
            if data_operation == "U":
                existing = get_supplier(
                    row_data.get("Shared Set", ""),
                    row_data.get("Supplier", ""),
                    tm.get(),
                )

                if not existing.get("data") or not existing["data"].get("supplierV2s"):
                    fail_count += 1
                    _mark_error(
                        ws, row_idx, status_col, error_col, header_row, [],
                        f"UPDATE failed: Supplier '{row_data.get('Supplier', '')}' not found in QAD "
                        f"(see debug log for the raw GET response)",
                    )
                    wb.save(file_path)
                    continue

                payload  = existing["data"]
                supplier = payload["supplierV2s"][0]

                supplier["businessRelationCode"]          = _cell_str(row_data, "Business Relation")
                supplier["isActive"]                      = _cell_str(row_data, "Active").lower() == "yes"
                supplier["supplierTypeCode"]              = _cell_str(row_data, "Supplier Type")
                supplier["purchaseTypeCode"]              = _cell_str(row_data, "Purchase Type")
                supplier["addressSearchName"]             = _cell_str(row_data, "Search Name")
                supplier["city"]                          = _cell_str(row_data, "City")
                supplier["stateCode"]                     = _cell_str(row_data, "State")
                supplier["street1"]                       = _cell_str(row_data, "Address 1")
                supplier["street2"]                       = _cell_str(row_data, "Address 2")
                supplier["zipCode"]                       = _cell_str(row_data, "Postal Code")
                supplier["telephone"]                     = _cell_str(row_data, "Telephone")
                supplier["EMail"]                         = _cell_str(row_data, "Email")
                supplier["currencyCode"]                  = _cell_str(row_data, "Currency")
                supplier["creditTermsCode"]               = _cell_str(row_data, "Credit Terms")
                supplier["invoiceStatusCode"]             = _cell_str(row_data, "Invoice Status")
                supplier["invoiceControlGLProfileCode"]   = _cell_str(row_data, "Invoice Control GL Profile")
                supplier["creditNoteControlGLProfileCode"]= _cell_str(row_data, "Credit Note Control GL Profile")
                supplier["prePaymentControlGLProfileCode"]= _cell_str(row_data, "Prepayment Control GL Profile")
                supplier["purchaseAccountGLProfileCode"]  = _cell_str(row_data, "Purchase Account GL Profile")
                supplier["taxZone"]                       = _cell_str(row_data, "Tax Zone")

            else:
                payload  = build_payload(row_data)
                supplier = payload["supplierV2s"][0]

            # -- Fold the bank entry (if any) into the same payload ------------------
            bank_ok, bank_err = _apply_bank_entry_to_supplier(supplier, row_data)
            if not bank_ok:
                fail_count += 1
                _log(f"row {row_idx} BANK ERROR -- {bank_err}")
                _mark_error(
                    ws, row_idx, status_col, error_col, header_row,
                    [OWN_BANK_NUMBER_COL], bank_err,
                )
                wb.save(file_path)
                continue

            has_bank_entry, _ = _bank_fields_present(row_data)
            expected_bank_entry = supplier["bankNumberRefV2s"][0] if has_bank_entry and supplier.get("bankNumberRefV2s") else None

            # -- Step 1: POST supplier (with one token-refresh retry on 401) ----------
            success   = False
            error_msg = ""

            for attempt in range(2):
                try:
                    success, error_msg = post_supplier(payload, tm.get(), is_create=is_create)
                    break
                except _TokenExpired:
                    if attempt == 0:
                        tm.refresh()
                        continue
                    error_msg = "Token refresh failed — unauthorised"
                    break
                except requests.RequestException as e:
                    error_msg = f"Network error: {e}"
                    _log(f"row {row_idx} NETWORK ERROR: {e}")
                    break

            if not success:
                fail_count += 1
                _mark_error(ws, row_idx, status_col, error_col, header_row, [], error_msg)
                wb.save(file_path)
                time.sleep(0.1)
                continue

            # -- Step 2 (CREATE only): update domain settings -------------------------
            if is_create:
                domain        = _cell_str(row_data, "Domain")
                site_code     = _cell_str(row_data, "Site Code")
                daybook_set   = _cell_str(row_data, "Daybook Set")
                supplier_code = _cell_str(row_data, "Supplier")

                domain_ok  = False
                domain_err = ""

                for attempt in range(2):
                    try:
                        domain_ok, domain_err = update_domain_settings(
                            supplier_code, domain, site_code, daybook_set, tm.get()
                        )
                        break
                    except _TokenExpired:
                        if attempt == 0:
                            tm.refresh()
                            continue
                        domain_err = "Token refresh failed — unauthorised"
                        break
                    except requests.RequestException as e:
                        domain_err = f"Network error: {e}"
                        break

                if not domain_ok:
                    fail_count += 1
                    _mark_error(
                        ws, row_idx, status_col, error_col, header_row, [],
                        f"Supplier created but domain settings failed: {domain_err}",
                    )
                    wb.save(file_path)
                    time.sleep(0.1)
                    continue

            # -- Step 3 (only if a bank entry was sent): verify it landed -------------
            if expected_bank_entry is not None:
                problems = []
                try:
                    problems = _verify_bank_entry(
                        row_data.get("Shared Set", ""),
                        row_data.get("Supplier", ""),
                        expected_bank_entry,
                        tm.get(),
                    )
                except _TokenExpired:
                    tm.refresh()
                    try:
                        problems = _verify_bank_entry(
                            row_data.get("Shared Set", ""),
                            row_data.get("Supplier", ""),
                            expected_bank_entry,
                            tm.get(),
                        )
                    except requests.RequestException as e:
                        problems = [f"Verification GET failed: {e}"]
                except requests.RequestException as e:
                    problems = [f"Verification GET failed: {e}"]

                if problems:
                    combined = "; ".join(problems)
                    _log(f"row {row_idx} VERIFY MISMATCH -- {combined}")
                    ok_count += 1
                    _mark_verify(ws, row_idx, status_col, error_col, combined)
                    wb.save(file_path)
                    time.sleep(0.1)
                    continue

            # -- All steps passed -------------------------------------------------------
            ok_count += 1
            _log(f"row {row_idx} DONE")
            _mark_done(ws, row_idx, status_col, error_col)
            wb.save(file_path)
            time.sleep(0.1)

        except Exception as e:
            fail_count += 1
            tb = traceback.format_exc()
            _log(f"UNEXPECTED EXCEPTION row {row_idx} "
                 f"(Supplier='{row_data.get('Supplier', '')}'):\n{tb}")
            _mark_error(
                ws, row_idx, status_col, error_col, header_row, [],
                f"Unexpected error: {e} (see debug log for full traceback)",
            )
            wb.save(file_path)
            continue

    return ok_count, fail_count


# =============================================================================
# 10. ORCHESTRATOR  (folder scan -> process every file -> return totals)
# =============================================================================

def run(folder_path: str) -> tuple[int, int]:
    """
    Scan folder_path for .xlsx files, process each one, return (total_ok, total_fail).
    This is the single entry point used by both main.py and __main__.
    """
    global _LOG

    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    tm = TokenManager()

    total_ok   = 0
    total_fail = 0

    for file_name in xlsx_files:
        file_path  = os.path.join(folder, file_name)
        file_label = os.path.splitext(file_name)[0]

        _LOG = DebugLog(folder, file_label)
        _LOG.write(f"Processing file: {file_path}")

        print(f"Loading : {file_path}")

        ok, fail   = process_file(file_path, tm)
        total_ok   += ok
        total_fail += fail

        print("\nSupplier Load Summary")
        print("-" * 40)
        print(f"  Rows loaded successfully : {ok}")
        print(f"  Rows failed              : {fail}")
        if fail == 0:
            print("  Result : ALL ROWS LOADED SUCCESSFULLY \u2713")
        else:
            print("  Result : COMPLETED WITH ERRORS — fix red rows and re-run")
        print("  NOTE: rows marked 'DONE (VERIFY)' in yellow saved successfully")
        print("        but post-write verification of the bank entry found a")
        print("        mismatch -- check the Error column and the debug log.")
        print()

        _LOG.write(f"File complete. ok={ok}, fail={fail}")
        print(f"  Debug log: {_LOG.path}\n")
        _LOG.close()
        _LOG = None

    print("=" * 40)
    print(f"  Total loaded : {total_ok}")
    print(f"  Total failed : {total_fail}")
    print("=" * 40)

    return total_ok, total_fail


# =============================================================================
# 11. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["supplier"])
    )
    ok, fail = run(folder)
    sys.exit(0 if fail == 0 else 1)