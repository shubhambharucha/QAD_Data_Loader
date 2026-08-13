"""
Customer_load.py
----------------
Loads customer rows from all .xlsx files in the configured customer folder
into QAD via the customerV2s API.

Behaviour
---------
- Uses the OAuth token already obtained at login (held server-side in
  main.py's SESSIONS store) — this script no longer fetches its own token
  from config, because config.json no longer contains service-account
  credentials (see TokenManager notes below).
- Skips rows with Status = DONE (for Create operations only)
- Lightweight mandatory-field check before any API call
- Success/error extraction normalized via qad_response_utils.normalize_response,
  which handles both response envelopes QAD has been observed sending:
    1. submitResult-wrapped (normal validation failures)
    2. bare top-level errors[] (gateway/permission failures, e.g. 403)
- Error DETAIL extracted from errors[] (fieldName, message, fieldValue, code)
  and translated to Excel column names via QAD_FIELD_TO_COLUMN /
  DOMAIN_FIELD_TO_COLUMN below, so the exact bad cell gets highlighted —
  not just Customer/Status/Error.
- On success  → clear all red fills, clear Error, set Status = DONE
- On failure  → red-fill first col + every bad col QAD identified + Error col,
  log the combined message, set Status = ERROR
- Processes every row regardless of individual failures
- Returns (ok_count, fail_count)

Generic Excel formatting (fills, mark DONE/ERROR, translating QAD's
field-level errors into highlighted columns) lives in excel_format_utils.py.
Generic response-envelope normalization lives in qad_response_utils.py.
Both are shared by every <Entity>_load.py loader script. This file only
owns entity-specific knowledge: which QAD field maps to which Excel
column, payload building, and API communication.

CREATE flow (Data Operation = C):
  1. POST to customerV2s  → creates customer
  2. GET  mfgCustomers    → fetch auto-created domain record
  3. PATCH + POST         → update siteCode + daybookSetCode
  If step 3 fails → row marked ERROR with domain-settings errors highlighted
  on the Site Code / Daybook Set / Domain columns where QAD identifies them.

UPDATE flow (Data Operation = U):
  1. GET  customerV2s     → fetch existing record
  2. Patch editable fields
  3. POST customerV2s     → update customer
  (Domain settings NOT touched on UPDATE — one-time setup only)
"""

import os
import sys
import time
import json
import requests
import openpyxl
from datetime import datetime

# ── Progress callback (used by main.py SSE streaming) ──────────────────

_progress_callback = None

def set_progress_callback(callback):
    global _progress_callback
    _progress_callback = callback


# ── Path / config ─────────────────────────────────────────────────────────
ROOT_DIR    = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT_DIR)
sys.path.insert(0, SCRIPTS_DIR)
from config import CONFIG
from excel_format_utils import RED_FILL, CLEAR_FILL, mark_done, mark_error, resolve_qad_errors
from qad_response_utils import normalize_response

# ── Mandatory columns (customer) ──────────────────────────────────────────
MANDATORY_COLUMNS = [
    "Customer",
    "Shared Set",
    "Business Relation",
    "Active",
    "Currency",
    "Credit Terms",
    "Invoice Status",
    "Invoice Control GL Profile",
    "Credit Note Control GL Profile",
    "Prepayment Control GL Profile",
    "Sales Account GL Profile",
]

# ── Mandatory columns for domain settings (CREATE only) ───────────────────
MANDATORY_DOMAIN_COLUMNS = [
    "Domain",
    "Site Code",
    "Daybook Set",
]

# ── QAD API fieldName → Excel column name ─────────────────────────────────
# Extend this map whenever a new QAD fieldName is observed in errors[]
# that should highlight a specific Excel cell. Anything NOT in this map still
# shows up in the Error message text — it just won't get its own highlighted
# cell (falls back to today's Customer/Status/Error-only highlighting).
#
# NOTE: QAD's key is "fieldName" (not "field") — resolve_qad_errors() reads
# fieldName. Previously this was misread and silently dropped every
# column-level highlight; fixed in excel_format_utils.py.
QAD_FIELD_TO_COLUMN = {
    "customerCode":                     "Customer",
    "sharedSetCode":                     "Shared Set",
    "businessRelationCode":              "Business Relation",
    "isActive":                          "Active",
    "currencyCode":                      "Currency",
    "customerCurrencyCode":              "Currency",
    "CreditTermsCode":                   "Credit Terms",
    "InvoiceStatusCode":                 "Invoice Status",
    "InvoiceControlGLProfileCode":       "Invoice Control GL Profile",
    "CreditNoteControlGLProfileCode":    "Credit Note Control GL Profile",
    "PrePaymentControlGLProfileCode":    "Prepayment Control GL Profile",
    "SalesAccountGLProfileCode":         "Sales Account GL Profile",
    "creditRatingCode":                  "Credit Rating",
    "TaxZone":                           "Tax Zone",
    "EMail":                             "Email",
    "telephone":                         "Telephone",
    "AddressSearchName":                 "Search Name",   # confirmed via QAD error log 2026-07-22

    # ── Below: audited against every val()-sourced field in build_payload().
    # Key names follow the same JSON-field convention as the confirmed
    # entries above. NOT yet individually confirmed against a live QAD
    # error (only "AddressSearchName" and the entries above it have been
    # seen for real) — QAD is known to sometimes use different casing
    # (addressSearchName in payload vs "AddressSearchName" in the error),
    # so if a highlight doesn't fire for one of these, check the real
    # fieldName in the debug print and correct the key here.
    "CustomerTypeCode":                  "Customer Type",   # confirmed via QAD error log 2026-07-22
    "businessRelationName":              "Name",
    "businessRelationName2":             "Second Name",
    "businessRelationName3":             "Third Name",
    "addressTypeCode":                   "Address Type",
    "street1":                           "Address 1",
    "street2":                           "Address 2",
    "street3":                           "Address 3",
    "city":                              "City",
    "zipCode":                           "Postal Code",
    "stateCode":                         "State",
    "stateTax":                          "State Tax",
    "countryCode":                       "Country",
    "languageCode":                      "Language",
    "deductionControlGLProfileCode":     "Deduction Control GL",
    "taxClass":                          "Tax Class",
    "taxUsage":                          "Tax Usage",
    "isTaxable":                         "Taxable(Yes/No)",
    "isTaxInCity":                       "Tax in City(Yes/No)",
    "isTaxIncluded":                     "Tax Included(Yes/No)",
    "isTaxReport":                       "Tax Report",
    "federalTax":                        "Federal Tax",
    "miscellaneousTax1":                 "Miscellaneous Tax 1",
    "miscellaneousTax2":                 "Miscellaneous Tax 2",
    "miscellaneousTax3":                 "Miscellaneous Tax 3",
    "fixedCreditLimit":                  "Fixed Credit Limit",
    "isFixedCreditLimit":                "Apply Fixed Ceiling",
    "isTurnOverCreditLimit":             "Apply % of Turnover",
    "isMaxDaysOverdueCreditLimit":       "Apply Maximum Days Overdue",
    "maxDaysCreditLimit":                "Maximum Days Overdue",
    "turnoverCreditLimitPercent":        "Percentage of Turnover",
    "isLockedCreditLimit":               "Credit Hold",
    "warningCreditLimitPercent":         "Warning Ceiling %",
    "creditAgencyReference":             "Credit Agency Ref",
    "IsOverruleAllowedSOCreditLimit":    "Overrule Allowed SO(Yes/No)",
    "IsCheckBeforeSOCreditLimit":        "Calculate before Order Entry(Yes/No)",
    "IsCheckAfterSOCreditLimit":         "Calculate after Order Entry(Yes/No)",
    "IsCheckBeforeInvoiceCreditLimit":   "Calculate before Invoice(Yes/No)",
    "IsCheckAfterInvoiceCreditLimit":    "Calculate after Invoice Entry(Yes/No)",
    "IsOverAllowedInvoiceCreditLimit":   "Overrule Allowed Invoice(Yes/No)",
    "IsIncludeDraftCreditLimit":         "Include Drafts(Yes/No)",
    "isIncludeOpenItemsCreditLimit":     "Include Open Items(Yes/No)",
    "isIncludeSOCheckCreditLimit":       "Include Sales Orders(Yes/No)",
    "customerIsInclDeduction":           "Include Deductions",
    "customCombo10":                     "Business Type",
}

# ── QAD domain-settings (mfgCustomers) fieldName → Excel column name ─────
DOMAIN_FIELD_TO_COLUMN = {
    "siteCode":       "Site Code",
    "daybookSetCode": "Daybook Set",
    "domainContext":  "Domain",
}


# =============================================================================
# 1. AUTHENTICATION
# =============================================================================
#
# CHANGED (see explanation in chat): this script used to fetch its own OAuth
# token from CONFIG['qad']['auth'] (a service-account username/password kept
# in config.json). config.json no longer has a "qad" key — it now has
# "environments": {"TEST": {...}, "PROD": {...}}, and those blocks only hold
# base_url/client_id/grant_type, NOT a username/password. The real
# username/password now only ever exists for the moment the user logs in via
# main.py's /api/login, which performs the OAuth exchange itself and keeps
# the resulting access_token server-side in SESSIONS[session_id].
#
# So TokenManager no longer knows how to mint its own token — it's handed
# the token (and the base_url that token is valid against) that main.py
# already obtained for this session, and just holds onto it for the run.
#
# If the token expires mid-run (401), this script CANNOT silently get a new
# one — it never had the password. TokenManager.refresh() raises instead of
# retrying, so process_file() surfaces a clear "please log in again" error
# on that row rather than the confusing bare KeyError('qad') seen before.

class TokenManager:
    """Holds the QAD access token + base_url for one run.

    token:    OAuth access token already obtained by main.py at login
              (stored server-side in SESSIONS[session_id]).
    base_url: the qracore base_url for the session's environment
              (same host used for the original OAuth call), e.g.
              CONFIG['environments']['TEST']['base_url'].
    """

    def __init__(self, token: str, base_url: str):
        if not token:
            raise ValueError("TokenManager requires a valid session access_token")
        if not base_url:
            raise ValueError("TokenManager requires a base_url")
        self._token    = token
        self.base_url  = base_url.rstrip("/")

    def get(self) -> str:
        return self._token

    def refresh(self) -> str:
        # No stored credentials to re-authenticate with from here — the
        # session's token came from the user's login and this script never
        # had the password. Surface this clearly instead of trying (and
        # failing) to hit an /oauth/token endpoint with nothing to send.
        raise RuntimeError(
            "QAD session token expired mid-run and cannot be refreshed "
            "automatically — please log in again and re-run the load."
        )


# =============================================================================
# 2. PAYLOAD BUILDER  (full payload — every field visible)
# =============================================================================

def build_payload(row: dict) -> dict:
    def val(col: str) -> str:
        v = row.get(col)
        return str(v).strip() if v is not None else ""

    def bool_val(col: str) -> bool:
        return val(col).lower() == "yes"

    def float_val(col: str) -> float:
        return float(val(col) or 0)

    def int_val(col: str) -> int:
        return int(float(val(col) or 0))

    def date_val(col: str) -> str:
        v = val(col)
        if not v:
            return ""
        try:
            # adjust input format based on source (excel date format)
            return datetime.strptime(v, "%d-%m-%Y").strftime("%Y-%m-%d")
        except Exception:
            return ""

    customer_code = val("Customer")
    shared_set    = val("Shared Set")
    uri           = f"urn:be:com.qad.base.customer.ICustomerV2:{shared_set}.{customer_code}"

    return {
        "supplementaryMessages": [],
        "customerV2s": [
            {
                # ── URIs / Identity ───────────────────────────────────────────
                "uri":                                      uri,
                "instanceURI":                              uri,
                "customerCode":                             customer_code,
                "customerID":                               0,
                "sharedSetCode":                            shared_set,
                "sharedSetID":                              0,

                # ── Business Relation ─────────────────────────────────────────
                "businessRelationCode":                     val("Business Relation"),
                "businessRelationID":                       0,
                "businessRelationName":                     val("Name"),
                "businessRelationName2":                    val("Second Name"),
                "businessRelationName3":                    val("Third Name"),
                "businessRelationConcurrencyHash":          "",
                "isBusinessRelationActive":                 True,
                "isBusinessRelationFieldsEnabled":          True,
                "isBusinessRelationIntercompany":           False,
                "isCreateBusinessRelationRequired":         True,
                "isInternalEntity":                         False,
                "intercompanyCode":                         "",

                # ── Address ───────────────────────────────────────────────────
                "addressID":                                0,
                "addressName":                              val("Name"),
                "addressSearchName":                        val("Search Name"),
                "addressTypeCode":                          val("Address Type"),

                "addressConcurrencyHash":                   "",
                "street1":                                  val("Address 1"),
                "street2":                                  val("Address 2"),
                "street3":                                  val("Address 3"),
                "city":                                     val("City"),
                "cityCode":                                 "",
                "zipCode":                                  val("Postal Code"),
                "stateCode":                                val("State"),
                "stateDescription":                         "",
                "stateTax":                                 val("State Tax"),
                "countryCode":                              val("Country"),
                "countryDescription":                       "",
                "countyCode":                               "",
                "countyDescription":                        "",
                "postalFormat":                             "0",
                "latitude":                                 0,
                "longitude":                                0,
                "isTemporaryAddress":                       False,

                # ── Status / Concurrency ──────────────────────────────────────
                "changeStatus":                             "2",
                "dataOperation":                            "",
                "concurrencyHash":                          "",
                "disallowedActions":                        "",
                "disallowedActionsMessage":                 "",
                "isPredefaulted":                           False,
                "isDomainRestricted":                       False,
                "lastModifiedDate":                         "",
                "lastModifiedTime":                         0,
                "lastModifiedUser":                         "",

                # ── Active / Type ─────────────────────────────────────────────
                "isActive":                                 bool_val("Active"),
                "customerTypeCode":                         val("Customer Type"),
                "customerTypeID":                           0,

                # ── Currency / Language ───────────────────────────────────────
                "currencyCode":                             val("Currency"),
                "currencyDescription":                      "",
                "currencyID":                               0,
                "customerCurrencyCode":                     val("Currency"),
                "languageCode":                             val("Language"),
                "languageDescription":                      "",

                # ── Credit Terms ──────────────────────────────────────────────
                "creditTermsCode":                          val("Credit Terms"),
                "creditTermsDescription":                   "",
                "creditTermsID":                            0,
                "creditTermsType":                          "",

                # ── Invoice / Status ──────────────────────────────────────────
                "invoiceStatusCode":                        val("Invoice Status"),
                "invoiceStatusDescription":                 "",
                "invoiceStatusID":                          0,
                "isInvoiceByAuthorization":                 False,
                "isWithPreInvoiceGroup":                    False,
                "isPrintBillWithItemDetails":               False,

                # ── GL Profiles ───────────────────────────────────────────────
                "invoiceControlGLProfileCode":              val("Invoice Control GL Profile"),
                "invoiceControlGLProfileDesc":              "",
                "invoiceControlGLProfileID":                0,
                "creditNoteControlGLProfileCode":           val("Credit Note Control GL Profile"),
                "creditNoteControlGLProfileDesc":           "",
                "creditNoteControlGLProfileID":             0,
                "prePaymentControlGLProfileCode":           val("Prepayment Control GL Profile"),
                "prePaymentControlGLProfileDesc":           "",
                "prePaymentControlGLProfileID":             0,
                "salesAccountGLProfileCode":                val("Sales Account GL Profile"),
                "salesAccountGLProfileDesc":                "",
                "salesAccountGLProfileID":                  0,
                "deductionControlGLProfileCode":            val("Deduction Control GL"),
                "deductionControlGLProfileDesc":            "",
                "deductionControlGLProfileID":              0,
                "financeChargeGLProfileCode":               "",
                "financeChargeGLProfileDesc":               "",
                "financeChargeGLProfileID":                 0,

                # ── Revenue Recognition GL Profiles ───────────────────────────
                "accruedRevRecGLProfileCode":               "",
                "accruedRevRecGLProfileDesc":               "",
                "accruedRevRecGLProfileID":                 0,
                "deferredRevRecGLProfileCode":              "",
                "deferredRevRecGLProfileDesc":              "",
                "deferredRevRecGLProfileID":                0,
                "COGSAccruedRevRecGLProfileCode":           "",
                "COGSAccruedRevRecGLProfileDesc":           "",
                "COGSAccruedRevRecGLProfileID":             0,
                "COGSDeferredRevRecGLProfileCode":          "",
                "COGSDeferredRevRecGLProfileDesc":          "",
                "COGSDeferredRevRecGLProfileID":            0,
                "COGSOffsetRevRecGLProfileCode":            "",
                "COGSOffsetRevRecGLProfileDesc":            "",
                "COGSOffsetRevRecGLProfileID":              0,
                "isAutoCreateRevRecContracts":              False,
                "revRecRuleID":                             0,
                "reviewRequiredForContracts":               "",

                # ── Tax ───────────────────────────────────────────────────────
                "taxZone":                                  val("Tax Zone"),
                "taxZoneDescription":                       "",
                "taxClass":                                 val("Tax Class"),
                "taxClassDescription":                      "",
                "taxDeclaration":                           0,
                "taxUsage":                                 val("Tax Usage"),
                "taxUsageDescription":                      "",
                "isTaxable":                                bool_val("Taxable(Yes/No)"),
                "isTaxInCity":                              bool_val("Tax in City(Yes/No)"),
                "isTaxIncluded":                             bool_val("Tax Included(Yes/No)"),
                "isTaxReport":                              bool_val("Tax Report"),
                "isLastFiling":                             False,
                "isReportedIN":                             False,
                "isElectronicInvoiceIN":                    False,
                "federalTax":                               val("Federal Tax"),
                "stateTax":                                 val("State Tax"),
                "miscellaneousTax1":                        val("Miscellaneous Tax 1"),
                "miscellaneousTax2":                        val("Miscellaneous Tax 2"),
                "miscellaneousTax3":                        val("Miscellaneous Tax 3"),
                "customerGTVatTransType":                   "",
                "vatDeliveryType":                          "",
                "vatPercentageLevel":                       "",
                "nameControl":                              "",
                "EORINumber":                                "",

                # ── Credit Limit ──────────────────────────────────────────────
                "fixedCreditLimit":                         float_val("Fixed Credit Limit"),
                "highCredit":                               0,
                "isFixedCreditLimit":                       bool_val("Apply Fixed Ceiling"),
                "isTurnOverCreditLimit":                    bool_val("Apply % of Turnover"),
                "isMaxDaysOverdueCreditLimit":              bool_val("Apply Maximum Days Overdue"),
                "maxDaysCreditLimit":                       int_val("Maximum Days Overdue"),
                "turnoverCreditLimitPercent":               float_val("Percentage of Turnover"),
                "isLockedCreditLimit":                      bool_val("Credit Hold"),
                "isToBeLockedCreditLimit":                  False,
                "warningCreditLimitPercent":                float_val("Warning Ceiling %"),
                "creditAgencyReference":                    val("Credit Agency Ref"),
                "creditRatingCode":                          val("Credit Rating"),
                "creditRatingID":                            0,

                # ── Credit Check ──────────────────────────────────────────────
                "isOverruleAllowedSOCreditLimit":           bool_val("Overrule Allowed SO(Yes/No)"),
                "isCheckBeforeSOCreditLimit":                bool_val("Calculate before Order Entry(Yes/No)"),
                "isCheckAfterSOCreditLimit":                 bool_val("Calculate after Order Entry(Yes/No)"),
                "isCheckBeforeInvoiceCreditLimit":          bool_val("Calculate before Invoice(Yes/No)"),
                "isCheckAfterInvoiceCreditLimit":           bool_val("Calculate after Invoice Entry(Yes/No)"),
                "isOverAllowedInvoiceCreditLimit":          bool_val("Overrule Allowed Invoice(Yes/No)"),
                "isIncludeDraftCreditLimit":                bool_val("Include Drafts(Yes/No)"),
                "isIncludeOpenItemsCreditLimit":            bool_val("Include Open Items(Yes/No)"),
                "isIncludeSOCheckCreditLimit":              bool_val("Include Sales Orders(Yes/No)"),
                "overToleranceAmount":                      0,
                "overTolerancePercent":                     0,
                "shortToleranceAmount":                     0,
                "shortTolerancePercent":                    0,
                "totalDaysLate":                            0,
                "totalNumberOfInvoices":                    0,

                # ── Finance Charges / Reminders / Statements ──────────────────
                "isFinanceCharge":                          False,
                "isPrintReminder":                          False,
                "isReminderRequired":                       False,
                "isPrintStatement":                         False,
                "reminderCountReset":                       False,
                "reminderAddressChangeStatus":              "",
                "reminderAddressID":                        0,
                "reminderAddressTypeCode":                  "",
                "reminderCustContactV2s":                   [],

                # ── Payment ───────────────────────────────────────────────────
                "paymentGroupCode":                         "",
                "paymentGroupDescription":                  "",
                "paymentGroupID":                           0,
                "domiciliationNumber":                      0,
                "isToleranceFromOwnBank":                   False,
                "isCompensationAllowed":                    False,

                # ── Billing ───────────────────────────────────────────────────
                "billToCustomerCode":                       "",
                "billToCustomerID":                         0,
                "billCollectorCode":                        "",
                "billCollectorID":                          0,
                "billingScheduleCode":                      "",
                "billingScheduleDescription":               "",
                "billingScheduleID":                        0,
                "statementCycle":                           "",
                "subAccountProfileCode":                    "",
                "subAccountProfileDesc":                    "",
                "subAccountProfileID":                      0,

                # ── Corporate Group ───────────────────────────────────────────
                "corporateGroupCode":                       "",
                "corporateGroupDescription":                "",
                "corporateGroupID":                         0,

                # ── Contact ───────────────────────────────────────────────────
                "EMail":                                    val("Email"),
                "fax":                                      "",
                "telephone":                                val("Telephone"),
                "webSite":                                  "",
                "commentNote":                              "",

                # ── Deduction ─────────────────────────────────────────────────
                "customerIsInclDeduction":                  bool_val("Include Deductions"),

                # ── Custom Fields ─────────────────────────────────────────────
                "customShort0":  "", "customShort1":  "", "customShort2":  "",
                "customShort3":  "", "customShort4":  "", "customShort5":  "",
                "customShort6":  "", "customShort7":  "", "customShort8":  "",
                "customShort9":  "", "customShort10": "", "customShort11": "",
                "customShort12": "", "customShort13": "", "customShort14": "",
                "customShort15": "", "customShort16": "", "customShort17": "",
                "customShort18": "", "customShort19": "",
                "customLong0":   "", "customLong1":   "",
                "customNote":    "",
                "customCombo0":  "", "customCombo1":  "", "customCombo2":  "",
                "customCombo3":  "", "customCombo4":  "", "customCombo5":  "",
                "customCombo6":  "", "customCombo7":  "", "customCombo8":  "",
                "customCombo9":  "", "customCombo10": val("Business Type"), "customCombo11": "",
                "customCombo12": "", "customCombo13": "", "customCombo14": "",
                "customDecimal0": 0, "customDecimal1": 0, "customDecimal2": 0,
                "customDecimal3": 0, "customDecimal4": 0,
                "customInteger0": 0, "customInteger1": 0, "customInteger2": 0,
                "customInteger3": 0, "customInteger4": 0,

                # ── Sub-lists ─────────────────────────────────────────────────
                "customerContactV2s":                       [],
                "customerSafDefaultV2s":                    [],
                "bankNumberRefV2s":                         [],
                "MDMBankNrSharedSets":                      [],
            }
        ],
    }


# =============================================================================
# 3. CUSTOMER API
# =============================================================================
#
# CHANGED: every function below used to build its URL from
# CONFIG['qad']['base_url']. That key no longer exists in config.json, so
# each function now takes an explicit base_url parameter instead — supplied
# by the caller (process_file), which gets it from tm.base_url. This keeps
# the functions themselves config-agnostic: they just talk to whatever host
# they're told to, which is also friendlier for TEST vs PROD since the same
# code path now naturally follows whichever environment the logged-in
# session belongs to.

class _TokenExpired(Exception):
    pass


def post_customer(payload: dict, token: str, base_url: str, is_create: bool = False) -> tuple[bool, list]:
    """
    POST customer payload to QAD.

    Returns (success, errors). errors is normalized via
    qad_response_utils.normalize_response(), so callers get a uniform
    list of dicts (fieldName/message/...) regardless of whether QAD sent
    back a submitResult-wrapped response or a bare top-level errors[]
    (e.g. 403 gateway/permission failures).

    Raises _TokenExpired on 401.
    """
    customer      = payload["customerV2s"][0]
    shared_set    = customer.get("sharedSetCode", "")
    customer_code = customer.get("customerCode", "")

    if is_create:
        url = (
            f"{base_url}/api/erp/customerV2s"
            f"?viewUri=urn:be:com.qad.base.customer.ICustomerV2"
        )
    else:
        url = (
            f"{base_url}/api/erp/customerV2s"
            f"?sharedSetCode={shared_set}&customerCode={customer_code}"
            f"&viewUri=urn:be:com.qad.base.customer.ICustomerV2"
        )

    resp = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type":  "application/json",
        },
        json=payload,
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    # ── TEMP DEBUG BLOCK — remove once field mapping is confirmed ─────────
    print("\n" + "="*70)
    print(f"DEBUG post_customer | HTTP {resp.status_code} | customer={customer_code}")
    print("="*70)
    print("RAW response text:")
    print(resp.text)
    print("-"*70)
    try:
        _debug_json = resp.json()
        print("Parsed JSON top-level keys:", list(_debug_json.keys()))
        print("Full parsed JSON:")
        print(json.dumps(_debug_json, indent=2))
    except Exception as _debug_exc:
        print(f"Could not parse response as JSON: {_debug_exc}")
    print("="*70 + "\n")
    # ── END TEMP DEBUG BLOCK ───────────────────────────────────────────────

    try:
        resp_json = resp.json()
    except Exception:
        resp_json = {}

    return normalize_response(resp_json, resp.status_code)


def get_customer(shared_set: str, customer_code: str, token: str, base_url: str) -> dict:
    url = f"{base_url}/api/erp/customerV2s"

    resp = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}"},
        params={
            "sharedSetCode": shared_set,
            "customerCode":  customer_code,
            "viewUri":       "urn:be:com.qad.base.customer.ICustomerV2",
        },
        timeout=30,
    )

    return resp.json()


# =============================================================================
# 4. DOMAIN SETTINGS API  (mfgCustomers)
# =============================================================================

def get_mfg_customer(domain: str, customer_code: str, token: str, base_url: str) -> dict:
    """Fetch the auto-created mfgCustomer domain record."""
    url = f"{base_url}/api/erp/mfgCustomers"

    resp = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}"},
        params={
            "domainContext": domain,
            "customerCode":  customer_code,
            "viewUri":       "urn:be:com.qad.base.customer.IMfgCustomer",
        },
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    return resp.json()


def post_mfg_customer(payload: dict, domain: str, customer_code: str, token: str, base_url: str) -> tuple[bool, list]:
    """
    POST updated domain settings back to QAD.

    Returns (success, errors) — same normalized shape as post_customer(),
    via qad_response_utils.normalize_response().
    """
    url = (
        f"{base_url}/api/erp/mfgCustomers"
        f"?domainContext={domain}&customerCode={customer_code}"
        f"&viewUri=urn:be:com.qad.base.customer.IMfgCustomer"
    )

    resp = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type":  "application/json",
        },
        json=payload,
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    try:
        resp_json = resp.json()
    except Exception:
        resp_json = {}

    # ── DEBUG BLOCK (remove before production) ────────────────────────────
    print("\n" + "="*60)
    print("DEBUG: POST mfgCustomers response")
    print("="*60)
    print(f"  HTTP Status   : {resp.status_code}")
    print(f"  Customer      : {customer_code}")
    print(f"  Domain        : {domain}")
    print("  Full response:")
    print(json.dumps(resp_json, indent=4))
    print("="*60 + "\n")
    # ── END DEBUG BLOCK ───────────────────────────────────────────────────

    return normalize_response(resp_json, resp.status_code)


def update_domain_settings(
    customer_code: str,
    domain:        str,
    site_code:     str,
    daybook_set:   str,
    token:         str,
    base_url:      str,
) -> tuple[bool, list]:
    """
    GET auto-created mfgCustomer → patch siteCode + daybookSetCode → POST back.
    Returns (success, errors) — errors is a normalized list, see post_mfg_customer().
    """
    existing = get_mfg_customer(domain, customer_code, token, base_url)

    mfg_list = (
        existing.get("data", {}).get("mfgCustomers")
        or existing.get("mfgCustomers")
    )

    if not mfg_list:
        return False, [{
            "fieldName": "domainContext",
            "message":   f"Domain settings record not found for customer '{customer_code}' in domain '{domain}'",
            "fieldValue": domain,
            "code":       None,
        }]

    payload    = existing.get("data", existing)
    mfg_record = mfg_list[0]

    mfg_record["siteCode"]       = site_code
    mfg_record["daybookSetCode"] = daybook_set

    return post_mfg_customer(payload, domain, customer_code, token, base_url)


# =============================================================================
# 5. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict) -> list[str]:
    missing = []
    for col in MANDATORY_COLUMNS:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


def _check_mandatory_domain(row_data: dict) -> list[str]:
    missing = []
    for col in MANDATORY_DOMAIN_COLUMNS:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 6. SINGLE-FILE PROCESSOR
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    """Open one workbook, process every row, return (ok_count, fail_count)."""

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # Row 1 = section labels, Row 2 = column headers, Row 3+ = data
    header_row = [
        str(c.value).strip() if c.value is not None else ""
        for c in ws[2]
    ]

    for col_name in ("Status", "Error"):
        if col_name not in header_row:
            ws.cell(row=2, column=len(header_row) + 1, value=col_name)
            header_row.append(col_name)

    status_col = header_row.index("Status") + 1
    error_col  = header_row.index("Error")  + 1

    ok_count   = 0
    fail_count = 0

    total_rows = ws.max_row - 2

    for row_idx, row in enumerate(ws.iter_rows(min_row=3), start=3):

        if _progress_callback:
            _progress_callback(row_idx - 2, total_rows)

        row_values = [cell.value for cell in row]

        if not any(row_values):
            continue

        row_data = dict(zip(header_row, row_values))

        status         = str(row_data.get("Status", "")).strip().upper()
        data_operation = str(row_data.get("Data Operation", "")).strip().upper() or "C"

        if status == "DONE" and data_operation == "C":
            continue

        # ── Mandatory check ───────────────────────────────────────────────
        missing = _check_mandatory(row_data)
        if missing:
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                missing,
                f"Missing mandatory fields: {', '.join(missing)}",
            )
            wb.save(file_path)
            continue

        # ── Operation validity check ──────────────────────────────────────
        if data_operation not in {"C", "U"}:
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                ["Data Operation"],
                "Data Operation must be C, U, or blank",
            )
            wb.save(file_path)
            continue

        # ── Domain settings mandatory check (CREATE only) ─────────────────
        is_create = (data_operation == "C")

        if is_create:
            missing_domain = _check_mandatory_domain(row_data)
            if missing_domain:
                fail_count += 1
                mark_error(
                    ws, row_idx, status_col, error_col, header_row,
                    missing_domain,
                    f"Missing domain setting fields: {', '.join(missing_domain)}",
                )
                wb.save(file_path)
                continue

        # ── Build payload ─────────────────────────────────────────────────
        if data_operation == "U":
            existing = get_customer(
                row_data.get("Shared Set", ""),
                row_data.get("Customer", ""),
                tm.get(),
                tm.base_url,
            )

            if not existing.get("data") or not existing["data"].get("customerV2s"):
                fail_count += 1
                mark_error(
                    ws, row_idx, status_col, error_col, header_row, [],
                    f"UPDATE failed: Customer '{row_data.get('Customer', '')}' not found in QAD",
                )
                wb.save(file_path)
                continue

            payload  = existing["data"]
            customer = payload["customerV2s"][0]

            def _patch_if_present(target_dict, key, raw_value, transform=lambda v: v):

                if raw_value is not None and str(raw_value).strip() != "":
                    target_dict[key] = transform(str(raw_value).strip())

            _patch_if_present(customer, "businessRelationCode",           row_data.get("Business Relation"))
            _patch_if_present(customer, "federalTax",                     row_data.get("Federal Tax"))
            _patch_if_present(customer, "isActive",                       row_data.get("Active"),   lambda v: v.lower() == "yes")
            _patch_if_present(customer, "currencyCode",                   row_data.get("Currency"))
            _patch_if_present(customer, "creditTermsCode",                row_data.get("Credit Terms"))
            _patch_if_present(customer, "creditRatingCode",                row_data.get("Credit Rating"))
            _patch_if_present(customer, "invoiceStatusCode",               row_data.get("Invoice Status"))
            _patch_if_present(customer, "taxZone",                         row_data.get("Tax Zone"))
            _patch_if_present(customer, "telephone",                       row_data.get("Telephone"))
            _patch_if_present(customer, "EMail",                            row_data.get("Email"))
            _patch_if_present(customer, "invoiceControlGLProfileCode",     row_data.get("Invoice Control GL Profile"))
            _patch_if_present(customer, "creditNoteControlGLProfileCode",  row_data.get("Credit Note Control GL Profile"))
            _patch_if_present(customer, "prePaymentControlGLProfileCode",  row_data.get("Prepayment Control GL Profile"))
            _patch_if_present(customer, "salesAccountGLProfileCode",       row_data.get("Sales Account GL Profile"))

        else:
            payload = build_payload(row_data)

        # ── Step 1: POST customer (with one token-refresh retry) ──────────
        success = False
        errors  = []

        for attempt in range(2):
            try:
                success, errors = post_customer(payload, tm.get(), tm.base_url, is_create=is_create)
                break
            except _TokenExpired:
                if attempt == 0:
                    try:
                        tm.refresh()
                        continue
                    except RuntimeError as refresh_exc:
                        errors = [{"fieldName": None, "message": str(refresh_exc), "fieldValue": None, "code": None}]
                        break
                errors = [{"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}]
                break
            except requests.RequestException as e:
                errors = [{"fieldName": None, "message": f"Network error: {e}", "fieldValue": None, "code": None}]
                break

        if not success:
            fail_count += 1
            bad_cols, error_msg = resolve_qad_errors(errors, QAD_FIELD_TO_COLUMN)
            mark_error(ws, row_idx, status_col, error_col, header_row, bad_cols, error_msg)
            wb.save(file_path)
            time.sleep(0.1)
            continue

        # ── Step 2 (CREATE only): update domain settings ──────────────────
        if is_create:
            domain        = str(row_data.get("Domain", "")).strip()
            site_code     = str(row_data.get("Site Code", "")).strip()
            daybook_set   = str(row_data.get("Daybook Set", "")).strip()
            customer_code = str(row_data.get("Customer", "")).strip()

            domain_ok     = False
            domain_errors = []

            for attempt in range(2):
                try:
                    domain_ok, domain_errors = update_domain_settings(
                        customer_code, domain, site_code, daybook_set, tm.get(), tm.base_url
                    )
                    break
                except _TokenExpired:
                    if attempt == 0:
                        try:
                            tm.refresh()
                            continue
                        except RuntimeError as refresh_exc:
                            domain_errors = [{"fieldName": None, "message": str(refresh_exc), "fieldValue": None, "code": None}]
                            break
                    domain_errors = [{"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}]
                    break
                except requests.RequestException as e:
                    domain_errors = [{"fieldName": None, "message": f"Network error: {e}", "fieldValue": None, "code": None}]
                    break

            if not domain_ok:
                fail_count += 1
                bad_cols, domain_msg = resolve_qad_errors(domain_errors, DOMAIN_FIELD_TO_COLUMN)
                mark_error(
                    ws, row_idx, status_col, error_col, header_row, bad_cols,
                    f"Customer created but domain settings failed: {domain_msg}",
                )
                wb.save(file_path)
                time.sleep(0.1)
                continue

        # ── All steps passed ──────────────────────────────────────────────
        ok_count += 1
        mark_done(ws, row_idx, status_col, error_col)
        wb.save(file_path)
        time.sleep(0.1)

    return ok_count, fail_count


# =============================================================================
# 7. ORCHESTRATOR
# =============================================================================
#
# CHANGED: run() is the standalone/CLI entry point (used when this script is
# executed directly, not through main.py's web backend). It has no session
# to borrow a token from, so it still needs to authenticate on its own. If
# you use this CLI path, set QAD_ENVIRONMENT / QAD_USERNAME / QAD_PASSWORD
# (or wire in your own prompt) so TokenManager can be constructed — see the
# __main__ block below. This keeps the standalone script usable without
# reintroducing a stored service-account password in config.json.

def run(folder_path: str, tm: TokenManager) -> tuple[int, int]:
    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    total_ok   = 0
    total_fail = 0

    for file_name in xlsx_files:
        file_path = os.path.join(folder, file_name)
        print(f"Loading : {file_path}")

        ok, fail   = process_file(file_path, tm)
        total_ok   += ok
        total_fail += fail

        print("\nLoad Summary")
        print("-" * 40)
        print(f"  Rows loaded successfully : {ok}")
        print(f"  Rows failed              : {fail}")
        if fail == 0:
            print("  Result : ALL ROWS LOADED SUCCESSFULLY ✓")
        else:
            print("  Result : COMPLETED WITH ERRORS — fix red rows and re-run")
        print()

    print("=" * 40)
    print(f"  Total loaded : {total_ok}")
    print(f"  Total failed : {total_fail}")
    print("=" * 40)

    return total_ok, total_fail


# =============================================================================
# 8. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    import getpass
    import requests as _requests

    environment = os.environ.get("QAD_ENVIRONMENT", "TEST").upper()
    env_cfg     = CONFIG["environments"][environment]

    username = os.environ.get("QAD_USERNAME") or input("QAD username: ")
    password = os.environ.get("QAD_PASSWORD") or getpass.getpass("QAD password: ")

    token_resp = _requests.post(
        f"{env_cfg['base_url']}/oauth/token",
        data={
            "client_id":  env_cfg["client_id"],
            "username":   username,
            "password":   password,
            "grant_type": env_cfg.get("grant_type", "password"),
        },
        timeout=30,
    )
    token_resp.raise_for_status()
    access_token = token_resp.json().get("access_token")
    if not access_token:
        sys.exit("ERROR: OAuth response did not contain access_token")

    tm = TokenManager(token=access_token, base_url=env_cfg["base_url"])

    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["customer"])
    )
    ok, fail = run(folder, tm)
    sys.exit(0 if fail == 0 else 1)