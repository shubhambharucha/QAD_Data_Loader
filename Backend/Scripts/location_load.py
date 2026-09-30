"""
location_load.py
-----------------
Loads Location Master records from all .xlsx files in the configured
Location folder into QAD via the locationMasters API
(urn:be:com.qad.inventory.inv.ILocationMaster).

Sheet layout
------------
One flat sheet (row 1 = headers, row 2+ = data):

    Site | Location | Description | Inventory Status | Location Type |
    Data Operation | Status | Error

CREATE only — same reasoning as site_load.py: only a single POST was ever
captured, so Update is not implemented rather than guessed at.

DOMAIN LOOKUP — why this loader makes an extra GET call
-------------------------------------------------------------------------
Location_sample.xlsx has no Domain column — only Site. But the captured
create call's URL requires domainCode as a query param
(...?domainCode=10USA&location=Dum01&siteCode=dummy&viewUri=...), and the
response body echoes domainCode too. A location can't be created without
one.

Rather than requiring the user to duplicate a Domain column here (or
guessing at a single fixed domain, which would silently misfile a
location the moment more than one domain is in play), this loader
resolves domainCode by doing one GET /sites?siteCode=... lookup per
unique Site value in the sheet (cached per run, not refetched per row —
same "cache per parent, not per row" approach as BOM_load.py's tree
lookups). This also gives a clean, specific row error ("Site 'X' not
found in QAD") if someone references a site that was never created,
instead of an opaque QAD error on the location call itself.

Auth
----
Same TokenManager(token, base_url) pattern as every other <Entity>_load.py
in this tool — no self-authentication, no stored credentials.

MINIMAL PAYLOAD — how it was derived
-------------------------------------------------------------------------
What was captured for /api/erp/locationMasters was the RESPONSE of a
successful create — a full echoed record including fields nobody set
(capacity, capacityUom, isReservedLocation, isSingleItem,
isSingleLotReference, isSupplierConsignmentEnabled, isTransferOwnership,
reservedLocations, projectCode, ...) plus read-only derived fields
(locationTypeDescription, inventoryStatusDescription, siteDescription —
these are QAD looking up the *Description* of a *Code* we sent, not
something to send ourselves). None of that is in the request payload
below; only fields that map to a Location_sample.xlsx column, plus the
same uri/concurrencyHash/disallowedActions/supplementaryMessages
boilerplate every other build_*_payload() in this codebase sends on
create.

ASSUMPTIONS (unconfirmed against a live sandbox — verify before trusting
in production):
  * "uri" is built as
    f"urn:be:com.qad.inventory.inv.ILocationMaster:{domain}.{site}.{location}",
    extending the {parent}.{code} convention one level deeper (domain ->
    site -> location). The captured response's uri
    ("urn:service:...ILocationMasterRead-LocationMasters:10USA.dummy.Dum01")
    is a READ-side service URI QAD generated on the way OUT, using a
    different view (ILocationMasterRead-LocationMasters, not the
    ILocationMaster business-entity view passed as viewUri on the
    request) — it was not copied verbatim, since it's not the shape a
    request uri would take.
  * "Inventory Status" is optional (the sample row leaves it blank) and
    is only included in the payload when the cell is non-empty AND isn't
    the literal value "default" — QAD presumably defaults it (the response
    showed "N-N-N") the same way it defaults Site's own
    defaultInventoryStatus. "default" is a dropdown option on the blank
    template (Scripts/blank_template.py) standing in for "leave it blank"
    — added because, per the requester, the manual Location Maintenance
    screen auto-fills this from the site's own default inventory status
    unless the user overrides it, so the dropdown needed an explicit way
    to say "use the site's default" rather than forcing a specific value.
  * dataOperation: "C" — same convention as site_load.py / po_load.py /
    so_load.py.
"""

import os
import sys
import time
import requests
import openpyxl

# ── Progress callback (used by main.py SSE streaming) ──────────────────

_progress_callback = None

def set_progress_callback(callback):
    global _progress_callback
    _progress_callback = callback


# ── Path / config ─────────────────────────────────────────────────────────
ROOT_DIR    = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT_DIR)
sys.path.insert(0, SCRIPTS_DIR)
from config import CONFIG
from excel_format_utils import mark_done, mark_error, resolve_qad_errors
from qad_response_utils import normalize_response

# ── Mandatory columns ────────────────────────────────────────────────────
# "Inventory Status" is deliberately NOT here — see module docstring.
MANDATORY_COLUMNS = [
    "Site",
    "Location",
    "Description",
    "Location Type",
]

# ── QAD API fieldName → Excel column name ───────────────────────────────
# domainCode is resolved from Site, not typed by the user, but a domainCode
# error from QAD is still mapped back to the Site column since that's the
# only lever the user has over it.
QAD_FIELD_TO_COLUMN = {
    "domainCode":          "Site",
    "siteCode":            "Site",
    "location":            "Location",
    "description":         "Description",
    "inventoryStatusCode": "Inventory Status",
    "locationType":        "Location Type",
}


# =============================================================================
# 1. AUTHENTICATION
# =============================================================================

class TokenManager:
    """Holds the QAD access token + base_url for one run.

    token:    OAuth access token already obtained by main.py at login
              (stored server-side in SESSIONS[session_id]).
    base_url: the qracore base_url for the session's environment
              (same host used for the original OAuth call), e.g.
              CONFIG['environments']['TEST']['base_url'].
    """

    def __init__(self, token: str, base_url: str):
        if not token:
            raise ValueError("TokenManager requires a valid session access_token")
        if not base_url:
            raise ValueError("TokenManager requires a base_url")
        self._token    = token
        self.base_url  = base_url.rstrip("/")

    def get(self) -> str:
        return self._token

    def refresh(self) -> str:
        # No stored credentials to re-authenticate with — the session's
        # token came from the user's login and this script never had the
        # password.
        raise RuntimeError(
            "QAD session token expired mid-run and cannot be refreshed "
            "automatically — please log in again and re-run the load."
        )


class _TokenExpired(Exception):
    pass


# =============================================================================
# 2. SITE LOOKUP  (domainCode resolution — see module docstring)
# =============================================================================

def get_site(site_code: str, token: str, base_url: str) -> dict:
    """
    GET /sites filtered by siteCode, to resolve the domainCode a location
    needs to be created under. Returns {} if not found.

    ASSUMPTION: the envelope shape for this GET was never captured (only
    the POST /sites response was). It's handled defensively below —
    tried as resp_json["data"]["sites"] first (the shape most GETs in
    this codebase use, e.g. Customer_load.get_customer()), falling back
    to a bare top-level resp_json["sites"] (the shape the POST /sites
    response actually used, per the captured trace) if "data" isn't
    present.

    Raises _TokenExpired on 401.
    """
    url = f"{base_url}/api/erp/sites"

    resp = requests.get(
        url,
        headers={"Authorization": f"Bearer {token}"},
        params={
            "siteCode": site_code,
            "viewUri":  "urn:be:com.qad.base.item.ISite",
        },
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    try:
        resp_json = resp.json()
    except Exception:
        return {}

    sites = (resp_json.get("data") or {}).get("sites") or resp_json.get("sites") or []
    return sites[0] if sites else {}


# =============================================================================
# 3. LOCATION API
# =============================================================================

def post_location(
    payload:   dict,
    domain:    str,
    site_code: str,
    location:  str,
    token:     str,
    base_url:  str,
) -> tuple[bool, list]:
    """
    POST locationMaster payload to QAD. Returns (success, errors),
    normalized via qad_response_utils.normalize_response().

    Raises _TokenExpired on 401.
    """
    url = (
        f"{base_url}/api/erp/locationMasters"
        f"?domainCode={domain}&location={location}&siteCode={site_code}"
        f"&viewUri=urn:be:com.qad.inventory.inv.ILocationMaster"
    )

    resp = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type":  "application/json",
        },
        json=payload,
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    try:
        resp_json = resp.json()
    except Exception:
        resp_json = {}

    return normalize_response(resp_json, resp.status_code)


# =============================================================================
# 4. VALUE EXTRACTORS
# =============================================================================

def sv(row_data: dict, key: str, default: str = "") -> str:
    val = row_data.get(key, default)
    return str(val).strip() if val is not None else default


# =============================================================================
# 5. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict) -> list[str]:
    missing = []
    for col in MANDATORY_COLUMNS:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 6. WORKBOOK HELPERS  (row 1 = headers)
# =============================================================================

def ensure_columns(ws, *col_names: str) -> list:
    """Add any missing column names to row 1 and return the full header list.

    Stops at the first blank cell — see site_load.py's ensure_columns() for
    why (blank_template.py's hidden dropdown-option columns further right
    would otherwise get read as if they were data columns).
    """
    header_row = []
    for cell in ws[1]:
        if cell.value is None:
            break
        header_row.append(cell.value.strip() if isinstance(cell.value, str) else cell.value)
    for name in col_names:
        if name not in header_row:
            ws.cell(row=1, column=len(header_row) + 1, value=name)
            header_row.append(name)
    return header_row


def rename_error(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if not name.startswith("error_"):
        new_path = os.path.join(folder, "error_" + name)
        os.rename(file_path, new_path)
        return new_path
    return file_path


def rename_restore(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if name.startswith("error_"):
        new_path = os.path.join(folder, name[len("error_"):])
        os.rename(file_path, new_path)
        return new_path
    return file_path


# =============================================================================
# 7. PAYLOAD BUILDER  (minimal — see module docstring)
# =============================================================================

def build_payload(row_data: dict, domain: str) -> dict:
    site_code   = sv(row_data, "Site")
    location    = sv(row_data, "Location")
    description = sv(row_data, "Description")
    inv_status  = sv(row_data, "Inventory Status")
    loc_type    = sv(row_data, "Location Type")

    uri = f"urn:be:com.qad.inventory.inv.ILocationMaster:{domain}.{site_code}.{location}"

    record = {
        "uri":                     uri,
        "domainCode":              domain,
        "siteCode":                site_code,
        "location":                location,
        "description":             description,
        "locationType":            loc_type,
        "dataOperation":           "C",
        "concurrencyHash":         "",
        "disallowedActions":       "",
        "disallowedActionsMessage": "",
    }

    # Optional — only sent when the user actually filled it in AND it isn't
    # the "default" sentinel. Leaving it blank, or explicitly choosing
    # "default" from the blank-template dropdown, means the same thing:
    # don't send inventoryStatusCode at all, so QAD inherits the site's own
    # default inventory status — matching what happens on the manual
    # Location Maintenance screen when this field is left untouched.
    if inv_status and inv_status.strip().lower() != "default":
        record["inventoryStatusCode"] = inv_status

    return {
        "supplementaryMessages": [],
        "locationMasters": [record],
    }


# =============================================================================
# 8. SINGLE-FILE PROCESSOR
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    print(f"\n{'─'*55}")
    print(f"  Loading: {os.path.basename(file_path)}")
    print(f"{'─'*55}")

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    header_row = ensure_columns(ws, "Status", "Error")
    status_col = header_row.index("Status") + 1
    error_col  = header_row.index("Error")  + 1

    ok_count   = 0
    fail_count = 0

    total_rows = ws.max_row - 1

    # Site -> domainCode, resolved at most once per unique site per run
    # (see module docstring). None = looked up and NOT found — cached too,
    # so a typo'd site doesn't trigger a fresh GET for every row using it.
    _site_domain_cache: dict[str, str | None] = {}

    def resolve_domain(site_code: str) -> tuple[str | None, list]:
        """Returns (domain_or_None, errors)."""
        if site_code in _site_domain_cache:
            cached = _site_domain_cache[site_code]
            if cached is None:
                return None, [{
                    "fieldName": "siteCode",
                    "message": f"Site '{site_code}' not found in QAD",
                    "fieldValue": site_code,
                    "code": None,
                }]
            return cached, []

        for attempt in range(2):
            try:
                site = get_site(site_code, tm.get(), tm.base_url)
                domain = str(site.get("domainCode", "")).strip() or None
                _site_domain_cache[site_code] = domain
                if domain is None:
                    return None, [{
                        "fieldName": "siteCode",
                        "message": f"Site '{site_code}' not found in QAD",
                        "fieldValue": site_code,
                        "code": None,
                    }]
                return domain, []
            except _TokenExpired:
                if attempt == 0:
                    try:
                        tm.refresh()
                        continue
                    except RuntimeError as refresh_exc:
                        return None, [{"fieldName": None, "message": str(refresh_exc), "fieldValue": None, "code": None}]
                return None, [{"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}]
            except requests.RequestException as e:
                return None, [{"fieldName": None, "message": f"Network error: {e}", "fieldValue": None, "code": None}]

        return None, [{"fieldName": None, "message": "Site lookup failed", "fieldValue": None, "code": None}]

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if _progress_callback:
            _progress_callback(row_idx - 1, total_rows)

        row_values = [cell.value for cell in row]
        if not any(row_values[:len(header_row)]):
            continue

        row_data = dict(zip(header_row, row_values))

        status         = str(row_data.get("Status", "")).strip().upper()
        data_operation = str(row_data.get("Data Operation", "")).strip().upper() or "C"

        if status == "DONE":
            print(f"  Row {row_idx} — already DONE, skipping")
            continue

        if data_operation != "C":
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                ["Data Operation"],
                "Only Data Operation = C (Create) is supported by this loader.",
            )
            wb.save(file_path)
            continue

        missing = _check_mandatory(row_data)
        if missing:
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                missing,
                f"Missing mandatory fields: {', '.join(missing)}",
            )
            wb.save(file_path)
            continue

        site_code = sv(row_data, "Site")
        location  = sv(row_data, "Location")
        print(f"  Location: {site_code} / {location}")

        domain, errors = resolve_domain(site_code)
        if domain is None:
            fail_count += 1
            bad_cols, error_msg = resolve_qad_errors(errors, QAD_FIELD_TO_COLUMN)
            mark_error(ws, row_idx, status_col, error_col, header_row, bad_cols, error_msg)
            print(f"  Location FAILED: {site_code}/{location} — {error_msg}")
            wb.save(file_path)
            continue

        payload = build_payload(row_data, domain)

        success = False
        errors  = []

        for attempt in range(2):
            try:
                success, errors = post_location(payload, domain, site_code, location, tm.get(), tm.base_url)
                break
            except _TokenExpired:
                if attempt == 0:
                    try:
                        tm.refresh()
                        continue
                    except RuntimeError as refresh_exc:
                        errors = [{"fieldName": None, "message": str(refresh_exc), "fieldValue": None, "code": None}]
                        break
                errors = [{"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}]
                break
            except requests.RequestException as e:
                errors = [{"fieldName": None, "message": f"Network error: {e}", "fieldValue": None, "code": None}]
                break

        if not success:
            fail_count += 1
            bad_cols, error_msg = resolve_qad_errors(errors, QAD_FIELD_TO_COLUMN)
            mark_error(ws, row_idx, status_col, error_col, header_row, bad_cols, error_msg)
            print(f"  Location FAILED: {site_code}/{location} — {error_msg}")
            wb.save(file_path)
            time.sleep(0.1)
            continue

        mark_done(ws, row_idx, status_col, error_col)
        ok_count += 1
        print(f"  Location OK: {site_code}/{location}")
        wb.save(file_path)
        time.sleep(0.1)

    print(f"\n  Locations — OK {ok_count} | Failed {fail_count}")

    return ok_count, fail_count


# =============================================================================
# 9. ORCHESTRATOR
# =============================================================================

def run(folder_path: str, tm: TokenManager) -> tuple[int, int]:
    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    total_success = 0
    total_fail    = 0

    for file_name in xlsx_files:
        file_path = os.path.join(folder, file_name)

        s, fail = process_file(file_path, tm)
        total_success += s
        total_fail    += fail

        if fail > 0:
            file_path = rename_error(file_path)
            print(f"\n  Errors found — renamed to: {os.path.basename(file_path)}")
        elif s > 0:
            file_path = rename_restore(file_path)

        print(f"\n  Load Summary {'-'*30}")
        print(f"    Rows loaded : {s}")
        print(f"    Rows failed : {fail}")
        if fail == 0:
            print("    Result      : ALL ROWS LOADED SUCCESSFULLY ✓")
        else:
            print("    Result      : COMPLETED WITH ERRORS — fix red rows and re-run")

    print(f"\n{'='*55}")
    print(f"  TOTAL — Success: {total_success} | Failed: {total_fail}")
    print(f"{'='*55}")

    return total_success, total_fail


# =============================================================================
# 10. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    import getpass
    import requests as _requests

    environment = os.environ.get("QAD_ENVIRONMENT", "TEST").upper()
    env_cfg     = CONFIG["environments"][environment]

    username = os.environ.get("QAD_USERNAME") or input("QAD username: ")
    password = os.environ.get("QAD_PASSWORD") or getpass.getpass("QAD password: ")

    token_resp = _requests.post(
        f"{env_cfg['base_url']}/oauth/token",
        data={
            "client_id":  env_cfg["client_id"],
            "username":   username,
            "password":   password,
            "grant_type": env_cfg.get("grant_type", "password"),
        },
        timeout=30,
    )
    token_resp.raise_for_status()
    access_token = token_resp.json().get("access_token")
    if not access_token:
        sys.exit("ERROR: OAuth response did not contain access_token")

    tm = TokenManager(token=access_token, base_url=env_cfg["base_url"])

    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["location"])
    )
    ok, fail = run(folder, tm)
    sys.exit(0 if fail == 0 else 1)