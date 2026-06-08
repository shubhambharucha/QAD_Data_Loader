import openpyxl
from openpyxl.styles import PatternFill

RED_FILL   = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
CLEAR_FILL = PatternFill(fill_type=None)

FIELD_RULES = {
    "Domain":     {"max_len": 8,  "type": "character"},
    "Field Name": {"max_len": 16, "type": "character"},
    "Value":      {"max_len": 20, "type": "character"},
    "Comment":    {"max_len": 60, "type": "character"},
    "Group":      {"max_len": 10, "type": "character"},
}

ENTITY_COL = "Domain"


def validate(file_path):
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    header_row = [cell.value for cell in ws[1]]

    if "Status" not in header_row:
        ws.cell(row=1, column=len(header_row) + 1, value="Status")
        header_row.append("Status")
    if "Error" not in header_row:
        ws.cell(row=1, column=len(header_row) + 1, value="Error")
        header_row.append("Error")

    status_col_idx = header_row.index("Status") + 1
    error_col_idx  = header_row.index("Error")  + 1
    entity_col_idx = header_row.index(ENTITY_COL) + 1

    has_errors  = False
    error_count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        row_values = [cell.value for cell in row]

        if not any(row_values):
            continue

        row_data = dict(zip(header_row, row_values))

        current_status = str(row_data.get("Status", "")).strip().upper()
        if current_status in ("DONE", "READY"):
            continue

        row_errors      = []
        error_cell_idxs = []

        for col_name, rule in FIELD_RULES.items():
            if col_name not in header_row:
                continue

            col_idx  = header_row.index(col_name) + 1
            value    = row_data.get(col_name)
            str_val  = str(value).strip() if value is not None else ""
            is_empty = str_val == "" or str_val.lower() == "none"

            if is_empty:
                row_errors.append(f"{col_name}: empty")
                error_cell_idxs.append(col_idx)
                continue

            max_len = rule.get("max_len")
            if max_len and len(str_val) > max_len:
                row_errors.append(f"{col_name}: max {max_len} chars (got {len(str_val)})")
                error_cell_idxs.append(col_idx)

        for cell in ws[row_idx]:
            cell.fill = CLEAR_FILL

        if row_errors:
            has_errors  = True
            error_count += 1
            error_msg   = "; ".join(row_errors)

            ws.cell(row=row_idx, column=status_col_idx, value="")
            ws.cell(row=row_idx, column=entity_col_idx).fill = RED_FILL
            for cidx in error_cell_idxs:
                ws.cell(row=row_idx, column=cidx).fill = RED_FILL
            ws.cell(row=row_idx, column=error_col_idx, value=error_msg).fill = RED_FILL

        else:
            ws.cell(row=row_idx, column=status_col_idx, value="READY")
            ws.cell(row=row_idx, column=error_col_idx,  value="")

    wb.save(file_path)
    return has_errors, error_count