import openpyxl
import requests
import time
import sys
import os
from datetime import datetime, timedelta
from openpyxl.styles import PatternFill

# ==========================================
# AUTH SETUP
# ==========================================
TOKEN_URL  = "https://cat5-devl.adaptive.qad.com/clouderp/oauth/token"
UPLOAD_URL = "https://cat5-devl.adaptive.qad.com/clouderp/api/erp/purchasing/supplierItems?viewUri=urn:be:com.qad.purchasing.setup.ISupplierItem"

AUTH_PARAMS = {
    "client_id": "afb97fd221925b87f01489aeb0e02e81",
    "username": "demo",
    "password": "qad",
    "grant_type": "password"
}

RED_FILL   = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
CLEAR_FILL = PatternFill(fill_type=None)


def get_new_token():
    try:
        response = requests.post(TOKEN_URL, params=AUTH_PARAMS)
        response.raise_for_status()
        token = response.json().get("access_token")
        if token:
            return token
        print("❌ Failed to obtain access token.")
        sys.exit(1)
    except Exception as e:
        print(f"❌ Token error: {e}")
        sys.exit(1)


def format_date(val):
    if isinstance(val, (int, float)):
        base = datetime(1899, 12, 30)
        return (base + timedelta(days=int(val))).strftime("%Y-%m-%dT00:00:00.000Z")
    elif isinstance(val, datetime):
        return val.strftime("%Y-%m-%dT00:00:00.000Z")
    return str(val).strip()


def ensure_columns(ws):
    header_row = [cell.value.strip() if isinstance(cell.value, str) else cell.value for cell in ws[1]]
    for name in ("Status", "Error"):
        if name not in header_row:
            ws.cell(row=1, column=len(header_row) + 1, value=name)
            header_row.append(name)
    return header_row, header_row.index("Status") + 1, header_row.index("Error") + 1


def mark_error(ws, row_idx, status_col, error_col, error_msg):
    ws.cell(row=row_idx, column=1).fill = RED_FILL
    ws.cell(row=row_idx, column=status_col, value="")
    ws.cell(row=row_idx, column=error_col, value=error_msg)
    ws.cell(row=row_idx, column=error_col).fill = RED_FILL


def mark_success(ws, row_idx, status_col, error_col):
    for cell in ws[row_idx]:
        cell.fill = CLEAR_FILL
    ws.cell(row=row_idx, column=status_col, value="DONE")
    ws.cell(row=row_idx, column=error_col, value="")


def rename_error(file_path):
    folder = os.path.dirname(file_path)
    name   = os.path.basename(file_path)
    if not name.startswith("error_"):
        new_path = os.path.join(folder, "error_" + name)
        os.rename(file_path, new_path)
        return new_path
    return file_path


def rename_restore(file_path):
    folder = os.path.dirname(file_path)
    name   = os.path.basename(file_path)
    if name.startswith("error_"):
        new_path = os.path.join(folder, name[len("error_"):])
        os.rename(file_path, new_path)
        return new_path
    return file_path


def parse_errors(resp_json):
    errors = resp_json.get("submitResult", {}).get("errors", [])
    msgs = []
    for e in errors:
        msg = e.get("message", "Unknown error")
        field = e.get("fieldName") or ""
        if "already exists" in msg.lower():
            msgs.append("Duplicate record")
        elif field:
            msgs.append(f"Field '{field}': {msg}")
        else:
            msgs.append(msg)
    return "; ".join(msgs) if msgs else "Failed"


def run(file_path):
    print(f"\n{'─'*55}")
    print(f"  📄 {os.path.basename(file_path)}")
    print(f"{'─'*55}")

    current_token = get_new_token()
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    header_row, status_col, error_col = ensure_columns(ws)

    success_count = skip_count = fail_count = 0

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        row_values = [cell.value for cell in row]
        if not any(row_values):
            continue

        row_data = dict(zip(header_row, row_values))

        if str(row_data.get("Status", "")).strip().upper() == "DONE":
            si = str(row_data.get("Supplier Item", "")).strip()
            print(f"  ⏭  {si or f'row {row_idx}'} — already DONE")
            skip_count += 1
            continue

        domain    = str(row_data.get("Domain Code", "")).strip()
        item      = str(row_data.get("Item Code", "")).strip()
        supplier  = str(row_data.get("Supplier Code", "")).strip()
        supp_item = str(row_data.get("Supplier Item", "")).strip()
        uom       = str(row_data.get("Unit of Measure", "")).strip()
        price     = row_data.get("Quote Price")
        date      = row_data.get("Quote Date")

        missing = []
        if not domain:    missing.append("Domain Code")
        if not item:      missing.append("Item Code")
        if not supplier:  missing.append("Supplier Code")
        if not supp_item: missing.append("Supplier Item")
        if not uom:       missing.append("Unit of Measure")
        if not price:     missing.append("Quote Price")
        if not date:      missing.append("Quote Date")

        if missing:
            msg = f"Missing required field(s): {', '.join(missing)}"
            mark_error(ws, row_idx, status_col, error_col, msg)
            print(f"  ✘  row {row_idx} — {msg}")
            fail_count += 1
            continue

        uri = f"urn:be:com.qad.purchasing.setup.ISupplierItem:{domain}.{item}.{supplier}.{supp_item}"

        item_payload = {
            "uri":           uri,
            "domainCode":    domain,
            "itemCode":      item,
            "supplierCode":  supplier,
            "supplierItem":  supp_item,
            "quoteDate":     format_date(date),
            "quotePrice":    price,
            "unitOfMeasure": uom,
            "dataOperation": "A",
            "currencyCode":  "INR"
        }

        pl = str(row_data.get("Price List", "")).strip()
        if pl:
            item_payload["priceList"] = pl

        payload = {"supplierItems": [item_payload]}

        try:
            response = requests.post(
                UPLOAD_URL, json=payload,
                headers={"Content-Type": "application/json", "Authorization": f"Bearer {current_token}"}
            )

            if response.status_code == 401:
                current_token = get_new_token()
                response = requests.post(
                    UPLOAD_URL, json=payload,
                    headers={"Content-Type": "application/json", "Authorization": f"Bearer {current_token}"}
                )

            resp_json = response.json()

            if response.status_code == 200 and resp_json.get("submitResult", {}).get("success"):
                mark_success(ws, row_idx, status_col, error_col)
                print(f"  ✔  Supplier Item: {supp_item} / Item: {item} / Supplier: {supplier}")
                success_count += 1
            else:
                error_msg = parse_errors(resp_json)
                mark_error(ws, row_idx, status_col, error_col, error_msg)
                print(f"  ✘  Supplier Item: {supp_item} — {error_msg}")
                fail_count += 1

        except Exception as e:
            mark_error(ws, row_idx, status_col, error_col, str(e))
            print(f"  ✘  Supplier Item: {supp_item} — Connection error: {e}")
            fail_count += 1

        time.sleep(0.1)

    wb.save(file_path)

    if fail_count > 0:
        file_path = rename_error(file_path)
        print(f"\n  ⚠  Errors found — file renamed to: {os.path.basename(file_path)}")
    elif success_count > 0 and fail_count == 0:
        file_path = rename_restore(file_path)

    print(f"\n  📊 Success: {success_count} | Skipped: {skip_count} | Failed: {fail_count}")
    return success_count, fail_count


# ==========================================
# ENTRY POINT
# ==========================================
if __name__ == "__main__":
    base   = os.path.dirname(os.path.abspath(__file__))
    folder = os.path.join(base, "..", "Supplier_Item")

    if not os.path.exists(folder):
        print(f"❌ Folder not found: {folder}")
        sys.exit(1)

    files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not files:
        print("⚠  No .xlsx files found in Supplier_Item/")
        sys.exit(0)

    total_s = total_f = 0
    for f in files:
        s, fail = run(os.path.join(folder, f))
        total_s += s
        total_f += fail

    print(f"\n{'═'*55}")
    print(f"  TOTAL — Success: {total_s} | Failed: {total_f}")
    print(f"{'═'*55}")
    sys.exit(0 if total_f == 0 else 1)