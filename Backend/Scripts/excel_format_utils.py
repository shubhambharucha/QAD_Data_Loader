"""
excel_format_utils.py
----------------------
Generic, entity-agnostic Excel formatting helpers shared by every
<Entity>_load.py loader script.

This module intentionally contains NO entity-specific knowledge (no field
names, no column lists). Each loader keeps its own field->column mapping
(e.g. Customer_load.py's QAD_FIELD_TO_COLUMN) and passes the results into
these generic functions.

Responsibilities:
- Fill constants (single source of truth for RED_FILL / CLEAR_FILL)
- Marking a row DONE
- Marking a row ERROR (identifier col + bad cols + Error col highlighted)
- Translating QAD's structured errors[] into (bad_columns, combined_message)
  using a caller-supplied field map

Nothing in here talks to QAD, opens files, or knows about main.py's SSE
stream. It only ever touches the openpyxl worksheet object it's handed.
"""

from openpyxl.styles import PatternFill

RED_FILL   = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
CLEAR_FILL = PatternFill(fill_type=None)


def mark_done(ws, row_idx: int, status_col: int, error_col: int) -> None:
    """Clear all fills on the row, set Status = DONE, clear Error."""
    for cell in ws[row_idx]:
        cell.fill = CLEAR_FILL
    ws.cell(row=row_idx, column=status_col, value="DONE")
    ws.cell(row=row_idx, column=error_col,  value="")


def mark_error(
    ws,
    row_idx:        int,
    status_col:     int,
    error_col:      int,
    header_row:     list,
    bad_columns:    list,
    error_msg:      str,
    identifier_col: int = 1,
) -> None:
    """
    Clear all fills, then highlight:
      - identifier_col (defaults to column 1, e.g. Customer)
      - status_col, error_col
      - every column name in bad_columns that exists in header_row

    Set Status = ERROR and write error_msg into the Error cell.
    """
    for cell in ws[row_idx]:
        cell.fill = CLEAR_FILL

    ws.cell(row=row_idx, column=identifier_col).fill = RED_FILL
    ws.cell(row=row_idx, column=status_col).fill      = RED_FILL
    ws.cell(row=row_idx, column=error_col).fill       = RED_FILL
    ws.cell(row=row_idx, column=status_col, value="ERROR")
    ws.cell(row=row_idx, column=error_col,  value=error_msg)

    for col_name in bad_columns:
        if col_name in header_row:
            ws.cell(row=row_idx, column=header_row.index(col_name) + 1).fill = RED_FILL


def resolve_qad_errors(errors: list, field_map: dict) -> tuple:
    """
    Translate QAD's structured errors[] into:
      - bad_columns  : list[str]  Excel column names to highlight
                       (deduplicated, order-preserved)
      - combined_msg : str        all error messages joined for the Error cell

    errors:
        list of dicts, straight from submitResult.errors[] or a bare
        top-level errors[] (see qad_response_utils.normalize_response),
        e.g.:
        [{"fieldName": "AddressSearchName", "message": "Field length
          should be less than 29.", "fieldValue": "...", "code": "47"}]

        NOTE: QAD's key is "fieldName", not "field" — this was previously
        misread, which silently dropped every column-level highlight.

    field_map:
        dict mapping QAD fieldName -> Excel column name, e.g.
        {"AddressSearchName": "Search Name", "currencyCode": "Currency"}

    Behavior:
        A message is ALWAYS included in combined_msg, even if its field has
        no entry in field_map (or fieldName is missing/empty) — no error
        information is ever silently dropped. Fields with no mapping don't
        get an extra cell highlight (falls back to identifier + Error only).

        If NONE of the errors carried a usable fieldName (e.g. a bare 403
        gateway/permission error with fieldName: ""), a short honest hint
        is appended instead of leaving the person with a dead-end message.
    """
    bad_columns  = []
    messages     = []

    for e in errors or []:
        msg   = (e.get("message") or "").strip()
        field = e.get("fieldName")

        if field and field in field_map:
            col_name = field_map[field]
            if col_name not in bad_columns:
                bad_columns.append(col_name)
            if msg:
                messages.append(f"{col_name}: {msg}")
        elif msg:
            code = e.get("code")
            messages.append(f"{code}: {msg}" if code else msg)

    combined_msg = "; ".join(messages) if messages else "QAD rejected the record (no message returned)"

    if not bad_columns:
        combined_msg += (
            " [no field identified by QAD — if this is a 403/access error, "
            "check Shared Set, Customer Code, or domain permissions first]"
        )

    return bad_columns, combined_msg