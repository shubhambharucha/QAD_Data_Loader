"""
receipts_unplanned_load.py
---------------------------
Loads Unplanned Receive transactions from all .xlsx files in the
configured folder into QAD via the unplannedReceives /
validateAndAssignDefaultsLine APIs
(urn:be:com.qad.logistics.unplanned.IUnplannedReceiveV2).

This is the "Receipts Unplanned" action reachable from the Inventory
Detail page's Actions menu.

Sheet layout
------------
One flat sheet (row 1 = headers, row 2+ = data) — ONE ROW = ONE COMPLETE
RECEIPT. Confirmed with the requester: despite the API's grid/header+lines
shape (unplannedProcessorHeaderV2s containing an
unplannedProcessorDetailV2s array), a single unplanned receive only ever
has exactly one detail line under its header — there is no "add another
line" flow in the real UI. So this is NOT a Header/Lines two-sheet loader
like po_load.py/so_load.py; it's a flat single-sheet loader like
site_load.py/location_load.py, just with more API steps per row:

    Domain | Site | Item Code | Quantity | Location | Lot Serial |
    Reference | Description | Data Operation | Status | Error

Only Domain, Site, Item Code and Quantity are mandatory — per the
requester, "we only need to enter item code, quantity, rest gets
handled" (itemCode triggers QAD to auto-populate credit/debit accounts,
inventory status, etc. from the item master — this loader never computes
those itself, same philosophy as every other entity here).

CREATE only — no Update flow was ever captured or asked for.

API sequence (captured via live browser trace, requester walked through
init -> fieldChange(quantity) -> fieldChange(itemCode) -> commit; order
below is init -> fieldChange(itemCode) -> fieldChange(quantity) -> commit,
since itemCode is what triggers the defaulting logic and there's no
reason to set quantity first for a loader starting from scratch, unlike
the manual UI session that captured this, which already had itemCode
pre-filled from the Inventory Detail page it was launched from):

  1. GET  /unplannedReceives                        -> blank header+detail
  2. POST /unplannedReceives/validateAndAssignDefaultsLine?fieldName=itemCode
         -> QAD assigns creditAcct/debitAcct/inventoryStatusCode/etc.
  3. POST /unplannedReceives/validateAndAssignDefaultsLine?fieldName=transactionQuantity
  4. POST /unplannedReceives                         -> commit (same URL as
         step 1, but POST, with the full accumulated header+detail tree
         as the body)

DEVIATIONS from the captured trace — both deliberate, both flagged here
so they're easy to revisit if QAD complains:

  * The captured GET init included siteCode/location/itemCode as query
    params, because that trace was launched FROM a specific Inventory
    Detail record that already had all three known. A bulk loader has no
    such page context to launch from, so GET init here only ever sends
    domainCode + siteCode (siteCode because it's mandatory and identifies
    "where" — same role po_load.py's init plays with domain+PO number).
    itemCode is always left blank at init time and set via its own
    fieldChange next, since fieldChange is what's actually named
    "validateAndAssignDefaultsLine" — the endpoint that runs the
    autopopulate logic init's query param was never confirmed to trigger.
  * Location / Lot Serial / Reference / Description (all optional) are
    NOT sent via their own fieldChange calls — none were ever captured,
    and unlike itemCode none of them drive derived defaults, so there's
    no evidence a dedicated fieldChange step is even needed for them.
    Instead they're overlaid directly onto the detail object returned by
    the itemCode fieldChange (same "overlay user values onto the
    template row" approach as BOM_load.py's apply_row_fields()) and ride
    along in the body of every POST from that point forward, all the way
    to commit.

MINIMAL PAYLOAD / ASSUMPTIONS (unconfirmed against a live sandbox —
verify before trusting in production):
  * Every capture available was the RESPONSE, never the request body, for
    both fieldChange calls and the commit call. The response shape is a
    nested tree: {"data": {"unplannedProcessorHeaderV2s": [{... header
    fields ..., "unplannedProcessorDetailV2s": [{... detail fields
    ...}]}]}} (GET init explicitly shows the "data" wrapper; the
    fieldChange/commit captures don't show it, but that's very likely
    just how far the requester expanded the DevTools tree when copying —
    every other entity in this codebase uses "data" consistently for GET
    vs POST). Response-parsing here tries resp_json["data"][...] first,
    falls back to a bare top-level key, matching the same defensive
    pattern location_load.py's get_site() already used for its own
    unconfirmed GET shape.
  * The request body for fieldChange/commit is assumed to be the FULL
    header object (not just the detail line in isolation, unlike
    po_load.py/so_load.py's flat purchaseOrderLines/salesOrderLines
    arrays) — i.e. {"unplannedProcessorHeaderV2s": [header]} where
    header["unplannedProcessorDetailV2s"] = [detail]. Reasoning: the
    header carries aggregate fields (totalCost, totalQuantity) that
    visibly changed after the transactionQuantity fieldChange in the
    captured response (totalQuantity went from 0 to 10) — QAD can only
    recompute those if it receives the header context, not just the
    isolated line.
  * submitResult success/error handling (qad_response_utils.normalize_
    response()) is assumed for fieldChange and commit responses, matching
    every other fieldChange/commit-style endpoint already confirmed
    working in this codebase (PO/SO/BOM/Customer/Site/Location) — the
    captured GET init response explicitly showed this envelope
    ("submitResult": {"success": true, "errors": [], ...}); the
    fieldChange/commit captures didn't show it, again likely just an
    un-expanded tree node rather than a different envelope.
  * "supplementaryMessages": [] at the top level of every POST body —
    boilerplate every other build_*_payload() in this codebase sends on
    create; not directly visible in what was captured (same "probably
    just not expanded" caveat as above) but harmless to include either
    way.

Auth
----
Same TokenManager(token, base_url) pattern as every other <Entity>_load.py
in this tool — no self-authentication, no stored credentials.
"""

import os
import sys
import time
import requests
import openpyxl

# ── Progress callback (used by main.py SSE streaming) ──────────────────

_progress_callback = None

def set_progress_callback(callback):
    global _progress_callback
    _progress_callback = callback


# ── Path / config ─────────────────────────────────────────────────────────
ROOT_DIR    = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT_DIR)
sys.path.insert(0, SCRIPTS_DIR)
from config import CONFIG
from excel_format_utils import mark_done, mark_error, resolve_qad_errors
from qad_response_utils import normalize_response

VIEW_URI = "urn:be:com.qad.logistics.unplanned.IUnplannedReceiveV2"

# ── Mandatory columns ────────────────────────────────────────────────────
MANDATORY_COLUMNS = [
    "Domain",
    "Site",
    "Item Code",
    "Quantity",
]

# ── QAD API fieldName → Excel column name ───────────────────────────────
QAD_FIELD_TO_COLUMN = {
    "domainCode":         "Domain",
    "siteCode":            "Site",
    "itemCode":             "Item Code",
    "transactionQuantity": "Quantity",
    "location":             "Location",
    "lotSerial":            "Lot Serial",
    "reference":            "Reference",
    "description":          "Description",
}


# =============================================================================
# 1. AUTHENTICATION
# =============================================================================

class TokenManager:
    """Holds the QAD access token + base_url for one run.

    token:    OAuth access token already obtained by main.py at login
              (stored server-side in SESSIONS[session_id]).
    base_url: the qracore base_url for the session's environment
              (same host used for the original OAuth call), e.g.
              CONFIG['environments']['TEST']['base_url'].
    """

    def __init__(self, token: str, base_url: str):
        if not token:
            raise ValueError("TokenManager requires a valid session access_token")
        if not base_url:
            raise ValueError("TokenManager requires a base_url")
        self._token    = token
        self.base_url  = base_url.rstrip("/")

    def get(self) -> str:
        return self._token

    def refresh(self) -> str:
        # No stored credentials to re-authenticate with — the session's
        # token came from the user's login and this script never had the
        # password.
        raise RuntimeError(
            "QAD session token expired mid-run and cannot be refreshed "
            "automatically — please log in again and re-run the load."
        )

    def headers(self) -> dict:
        return {
            "Content-Type":  "application/json",
            "Authorization": f"Bearer {self.get()}",
        }


# =============================================================================
# 2. HTTP HELPERS  (same "refresh once on 401, no retry once refresh itself
#    fails" pattern as po_load.py / so_load.py)
# =============================================================================

def api_get(url: str, tm: TokenManager) -> tuple[requests.Response, dict]:
    resp = None
    for attempt in range(2):
        resp = requests.get(url, headers=tm.headers(), timeout=30)
        if resp.status_code == 401 and attempt == 0:
            try:
                tm.refresh()
            except RuntimeError:
                break
            continue
        try:
            return resp, resp.json()
        except Exception:
            return resp, {}
    try:
        return resp, resp.json()
    except Exception:
        return resp, {}


def api_post(url: str, payload: dict, tm: TokenManager) -> tuple[requests.Response, dict]:
    resp = None
    for attempt in range(2):
        resp = requests.post(url, json=payload, headers=tm.headers(), timeout=30)
        if resp.status_code == 401 and attempt == 0:
            try:
                tm.refresh()
            except RuntimeError:
                break
            continue
        try:
            return resp, resp.json()
        except Exception:
            return resp, {}
    try:
        return resp, resp.json()
    except Exception:
        return resp, {}


def _network_error(exc: Exception) -> dict:
    return {"fieldName": None, "message": f"Network error: {exc}", "fieldValue": None, "code": None}


def _token_error() -> dict:
    return {
        "fieldName": None,
        "message": "QAD session token expired or invalid — please log in again and re-run the load.",
        "fieldValue": None,
        "code": None,
    }


def _http_error(status_code: int) -> dict:
    return {"fieldName": None, "message": f"Unexpected response — HTTP {status_code}", "fieldValue": None, "code": None}


def safe_get(url: str, tm: TokenManager) -> tuple[requests.Response | None, dict | None, list | None]:
    """GET wrapped with network-error trapping. Not run through
    normalize_response() — GET responses here are plain data payloads,
    not submit envelopes (mirrors po_load.py's safe_get())."""
    try:
        resp, resp_json = api_get(url, tm)
    except requests.RequestException as e:
        return None, None, [_network_error(e)]

    if resp.status_code == 401:
        return resp, resp_json, [_token_error()]
    if resp.status_code != 200:
        return resp, resp_json, [_http_error(resp.status_code)]
    return resp, resp_json, None


def safe_post(url: str, payload: dict, tm: TokenManager) -> tuple[requests.Response | None, dict | None, list | None]:
    """POST wrapped with network-error trapping and normalize_response()
    (see module docstring's ASSUMPTIONS on why that's expected to apply
    here too)."""
    try:
        resp, resp_json = api_post(url, payload, tm)
    except requests.RequestException as e:
        return None, None, [_network_error(e)]

    if resp.status_code == 401:
        return resp, resp_json, [_token_error()]

    success, errors = normalize_response(resp_json, resp.status_code)
    if not success:
        if not errors:
            errors = [_http_error(resp.status_code)]
        return resp, resp_json, errors
    return resp, resp_json, None


# =============================================================================
# 3. API — URL builders + response-shape extraction
# =============================================================================
# See module docstring's "DEVIATIONS" note for why init deliberately omits
# location/itemCode/lotSerial/reference, unlike the captured trace.

def _init_url(base_url: str, domain: str, site: str) -> str:
    return (
        f"{base_url}/api/erp/unplannedReceives"
        f"?domainCode={domain}&siteCode={site}&location=&itemCode=&lotSerial=&reference="
        f"&viewUri={VIEW_URI}"
    )


def _field_change_url(base_url: str, domain: str, header_id, detail_id, field_name: str) -> str:
    return (
        f"{base_url}/api/erp/unplannedReceives/validateAndAssignDefaultsLine"
        f"?domainCode={domain}&unplannedProcessorHeaderID={header_id}"
        f"&unplannedProcessorDetailID={detail_id}&fieldName={field_name}"
    )


def _commit_url(base_url: str, domain: str) -> str:
    # Matches the captured commit call exactly: only domainCode populated,
    # every other query param present but blank.
    return (
        f"{base_url}/api/erp/unplannedReceives"
        f"?domainCode={domain}&siteCode=&location=&itemCode=&lotSerial=&reference="
        f"&viewUri={VIEW_URI}"
    )


def _extract_header(resp_json: dict) -> dict:
    """Pull the single header object (with its single nested detail line)
    out of a response, trying the "data"-wrapped shape first (confirmed
    for GET init) and falling back to a bare top-level key (what the
    fieldChange/commit captures showed) — see module docstring."""
    headers = (
        (resp_json.get("data") or {}).get("unplannedProcessorHeaderV2s")
        or resp_json.get("unplannedProcessorHeaderV2s")
        or []
    )
    return headers[0] if headers else {}


def _extract_detail(header: dict) -> dict:
    details = header.get("unplannedProcessorDetailV2s") or []
    return details[0] if details else {}


# =============================================================================
# 4. VALUE EXTRACTORS
# =============================================================================

def sv(row_data: dict, key: str, default: str = "") -> str:
    val = row_data.get(key, default)
    return str(val).strip() if val is not None else default


def fv(row_data: dict, key: str, default: float = 0.0) -> float:
    try:
        return float(row_data.get(key, default))
    except (TypeError, ValueError):
        return default


# =============================================================================
# 5. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict) -> list[str]:
    missing = []
    for col in MANDATORY_COLUMNS:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 6. WORKBOOK HELPERS  (row 1 = headers)
# =============================================================================

def ensure_columns(ws, *col_names: str) -> list:
    """Add any missing column names to row 1 and return the full header list.

    Stops at the first blank cell — real headers are always contiguous
    from column 1, so this naturally excludes blank_template.py's hidden
    dropdown-option columns further right from being read as if they were
    data columns (see site_load.py's ensure_columns() for the full story).
    """
    header_row = []
    for cell in ws[1]:
        if cell.value is None:
            break
        header_row.append(cell.value.strip() if isinstance(cell.value, str) else cell.value)
    for name in col_names:
        if name not in header_row:
            ws.cell(row=1, column=len(header_row) + 1, value=name)
            header_row.append(name)
    return header_row


def rename_error(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if not name.startswith("error_"):
        new_path = os.path.join(folder, "error_" + name)
        os.rename(file_path, new_path)
        return new_path
    return file_path


def rename_restore(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if name.startswith("error_"):
        new_path = os.path.join(folder, name[len("error_"):])
        os.rename(file_path, new_path)
        return new_path
    return file_path


# =============================================================================
# 7. RECEIPT CREATION FLOW
# =============================================================================

def create_receipt(row_data: dict, tm: TokenManager) -> tuple[bool, list]:
    """
    Full multi-step create: init -> fieldChange(itemCode) -> overlay
    optional fields -> fieldChange(transactionQuantity) -> commit.

    Returns (success, errors) — errors is a normalized list of
    {"fieldName", "message", "fieldValue", "code"} dicts (empty on
    success) ready for resolve_qad_errors().
    """
    domain    = sv(row_data, "Domain")
    site      = sv(row_data, "Site")
    item_code = sv(row_data, "Item Code")
    qty       = fv(row_data, "Quantity")
    location  = sv(row_data, "Location")
    lot       = sv(row_data, "Lot Serial")
    reference = sv(row_data, "Reference")
    description = sv(row_data, "Description")

    # 1. Init — blank header+detail, QAD assigns headerID/detailID
    resp, resp_json, errors = safe_get(_init_url(tm.base_url, domain, site), tm)
    if errors:
        return False, errors

    header = _extract_header(resp_json or {})
    detail = _extract_detail(header)
    if not header or not detail:
        return False, [{"fieldName": None, "message": "Init returned no header/detail object", "fieldValue": None, "code": None}]

    header_id = header.get("unplannedProcessorHeaderID")
    detail_id = detail.get("unplannedProcessorDetailID")

    # 2. fieldChange: itemCode — the step that actually triggers QAD's
    #    autopopulate (credit/debit accounts, inventoryStatusCode, ...)
    detail["itemCode"] = item_code
    header["unplannedProcessorDetailV2s"] = [detail]
    resp, resp_json, errors = safe_post(
        _field_change_url(tm.base_url, domain, header_id, detail_id, "itemCode"),
        {"supplementaryMessages": [], "unplannedProcessorHeaderV2s": [header]}, tm,
    )
    if errors:
        return False, errors
    new_header = _extract_header(resp_json or {})
    if not new_header:
        return False, [{"fieldName": "itemCode", "message": "fieldChange(itemCode) returned no header", "fieldValue": item_code, "code": None}]
    header = new_header
    detail = _extract_detail(header)

    # 3. Overlay optional fields directly — no dedicated fieldChange call
    #    exists/was captured for any of these (see module docstring).
    if location:
        detail["location"] = location
    if lot:
        detail["lotSerial"] = lot
    if reference:
        detail["reference"] = reference
    if description:
        detail["description"] = description
    header["unplannedProcessorDetailV2s"] = [detail]

    # 4. fieldChange: transactionQuantity
    detail["transactionQuantity"] = qty
    header["unplannedProcessorDetailV2s"] = [detail]
    resp, resp_json, errors = safe_post(
        _field_change_url(tm.base_url, domain, header_id, detail_id, "transactionQuantity"),
        {"supplementaryMessages": [], "unplannedProcessorHeaderV2s": [header]}, tm,
    )
    if errors:
        return False, errors
    new_header = _extract_header(resp_json or {})
    if not new_header:
        return False, [{"fieldName": "transactionQuantity", "message": "fieldChange(transactionQuantity) returned no header", "fieldValue": qty, "code": None}]
    header = new_header

    # 5. Commit
    resp, resp_json, errors = safe_post(_commit_url(tm.base_url, domain), {"supplementaryMessages": [], "unplannedProcessorHeaderV2s": [header]}, tm)
    if errors:
        return False, errors

    return True, []


# =============================================================================
# 8. SINGLE-FILE PROCESSOR
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    print(f"\n{'─'*55}")
    print(f"  Loading: {os.path.basename(file_path)}")
    print(f"{'─'*55}")

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    header_row = ensure_columns(ws, "Status", "Error")
    status_col = header_row.index("Status") + 1
    error_col  = header_row.index("Error")  + 1

    ok_count   = 0
    fail_count = 0

    total_rows = ws.max_row - 1

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if _progress_callback:
            _progress_callback(row_idx - 1, total_rows)

        row_values = [cell.value for cell in row]
        if not any(row_values[:len(header_row)]):
            continue

        row_data = dict(zip(header_row, row_values))

        status         = str(row_data.get("Status", "")).strip().upper()
        data_operation = str(row_data.get("Data Operation", "")).strip().upper() or "C"

        if status == "DONE":
            print(f"  Row {row_idx} — already DONE, skipping")
            continue

        if data_operation != "C":
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                ["Data Operation"],
                "Only Data Operation = C (Create) is supported by this loader.",
            )
            wb.save(file_path)
            continue

        missing = _check_mandatory(row_data)
        if missing:
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                missing,
                f"Missing mandatory fields: {', '.join(missing)}",
            )
            wb.save(file_path)
            continue

        item_code = sv(row_data, "Item Code")
        qty       = fv(row_data, "Quantity")
        print(f"  Unplanned Receive: {item_code} | qty={qty}")

        success, errors = create_receipt(row_data, tm)

        if not success:
            fail_count += 1
            bad_cols, error_msg = resolve_qad_errors(errors, QAD_FIELD_TO_COLUMN)
            mark_error(ws, row_idx, status_col, error_col, header_row, bad_cols, error_msg)
            print(f"  Receipt FAILED: {item_code} — {error_msg}")
            wb.save(file_path)
            time.sleep(0.1)
            continue

        mark_done(ws, row_idx, status_col, error_col)
        ok_count += 1
        print(f"  Receipt OK: {item_code}")
        wb.save(file_path)
        time.sleep(0.1)

    print(f"\n  Unplanned Receives — OK {ok_count} | Failed {fail_count}")

    return ok_count, fail_count


# =============================================================================
# 9. ORCHESTRATOR
# =============================================================================

def run(folder_path: str, tm: TokenManager) -> tuple[int, int]:
    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    total_success = 0
    total_fail    = 0

    for file_name in xlsx_files:
        file_path = os.path.join(folder, file_name)

        s, fail = process_file(file_path, tm)
        total_success += s
        total_fail    += fail

        if fail > 0:
            file_path = rename_error(file_path)
            print(f"\n  Errors found — renamed to: {os.path.basename(file_path)}")
        elif s > 0:
            file_path = rename_restore(file_path)

        print(f"\n  Load Summary {'-'*30}")
        print(f"    Rows loaded : {s}")
        print(f"    Rows failed : {fail}")
        if fail == 0:
            print("    Result      : ALL ROWS LOADED SUCCESSFULLY ✓")
        else:
            print("    Result      : COMPLETED WITH ERRORS — fix red rows and re-run")

    print(f"\n{'='*55}")
    print(f"  TOTAL — Success: {total_success} | Failed: {total_fail}")
    print(f"{'='*55}")

    return total_success, total_fail


# =============================================================================
# 10. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    import getpass
    import requests as _requests

    environment = os.environ.get("QAD_ENVIRONMENT", "TEST").upper()
    env_cfg     = CONFIG["environments"][environment]

    username = os.environ.get("QAD_USERNAME") or input("QAD username: ")
    password = os.environ.get("QAD_PASSWORD") or getpass.getpass("QAD password: ")

    token_resp = _requests.post(
        f"{env_cfg['base_url']}/oauth/token",
        data={
            "client_id":  env_cfg["client_id"],
            "username":   username,
            "password":   password,
            "grant_type": env_cfg.get("grant_type", "password"),
        },
        timeout=30,
    )
    token_resp.raise_for_status()
    access_token = token_resp.json().get("access_token")
    if not access_token:
        sys.exit("ERROR: OAuth response did not contain access_token")

    tm = TokenManager(token=access_token, base_url=env_cfg["base_url"])

    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["receipts_unplanned"])
    )
    ok, fail = run(folder, tm)
    sys.exit(0 if fail == 0 else 1)