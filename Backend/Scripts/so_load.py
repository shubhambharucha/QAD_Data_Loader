"""
SO_load.py
----------
Loads Sales Order headers and lines from all .xlsx files in the
configured SalesOrder folder into QAD via the salesOrderHeaders /
salesOrderLines APIs.

Behaviour
---------
- Reads base_url and auth from config.json (via config.py) — no more
  hardcoded BASE_URL / AUTH_PARAMS.
- Fetches one OAuth token per run via a TokenManager; refreshes
  automatically on 401 (api_get()/api_post() retry once internally —
  same pattern as PO_load.py / Customer_load.py).
- Skips rows with Status = DONE.
- Error/response handling now shares the same framework as
  Customer_load.py / PO_load.py:
    * qad_response_utils.normalize_response() turns every QAD POST
      response into a uniform (success, errors[]) shape, regardless of
      whether QAD sent a submitResult-wrapped envelope or a bare
      top-level errors[] (gateway/permission failures).
    * excel_format_utils.resolve_qad_errors() maps each errors[]
      fieldName back to the Excel column name via
      SO_HEADER_FIELD_TO_COLUMN / SO_LINE_FIELD_TO_COLUMN below, so the
      exact bad cell gets highlighted instead of just col 1 + Error.
    * excel_format_utils.mark_done() / mark_error() replace the old
      mark_row() helper and share the same RED_FILL / CLEAR_FILL
      constants used by every other loader.
- Lightweight mandatory-field check runs before any API call, for both
  Header rows and Line rows.
- Network errors, persistent-401 (token refresh failure), and gateway/
  permission errors are all surfaced as normalized error dicts so they
  flow through the exact same resolve_qad_errors() / mark_error() path
  as ordinary field validation errors.
- Saves workbook after every row so progress survives a crash.
- Renames file to error_<name> if any failures occurred, restores the
  name if a previously-errored file now loads clean.
- Returns (total_success, total_fail).

Sheet layout (UNCHANGED)
------------------------
Workbook must contain two sheets: 'Header' and 'Lines'.
Each sheet must have a header row (row 1) with column names.

Sales Order business logic / API workflow — UNCHANGED
-------------------------------------------------------
  Header : single POST to salesOrderHeaders
  Lines  : GET  init (line number comes from the Excel "Line Number"
                 column if present, else sequential position — this
                 is SO-specific and intentionally different from the
                 PO loader, which lets QAD assign the line number)
           POST fieldChange(itemCode)
           POST fieldChange(siteCode)
           POST fieldChange(quantityOrdered)
           POST fieldChange(listPrice)          — only if > 0
           POST fieldChange(discountFormatted)  — only if > 0
           POST fieldChange(netPrice)           — only if > 0
           POST fieldChange(dueDate)
           POST sync to salesOrderLinesGrid
None of the endpoints, payload fields, sequencing, conditional price/
discount/net steps, or line-number handling have been changed — only
how auth/config is sourced and how responses are parsed / Excel is
annotated.
"""

import os
import sys
import time
import requests
import openpyxl
from datetime import datetime, timezone, timedelta
from collections import defaultdict

# ── Progress callback (used by main.py SSE streaming) ──────────────────

_progress_callback = None

def set_progress_callback(callback):
    global _progress_callback
    _progress_callback = callback


# ── Path / config ─────────────────────────────────────────────────────────
ROOT_DIR    = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT_DIR)
sys.path.insert(0, SCRIPTS_DIR)
from config import CONFIG
from excel_format_utils import RED_FILL, CLEAR_FILL, mark_done, mark_error, resolve_qad_errors
from qad_response_utils import normalize_response

_BASE_URL = CONFIG["qad"]["base_url"]

# ── API endpoint templates (UNCHANGED) ─────────────────────────────────────
HEADER_CREATE_URL = f"{_BASE_URL}/api/erp/salesOrderHeaders?viewUri=urn:be:com.qad.sales.salesorder.ISalesOrderHeader"
INIT_LINE_URL     = f"{_BASE_URL}/api/erp/salesOrderLinesGrid?initialize=true&domainCode={{domain}}&salesOrderNumber={{so}}"
FIELD_CHANGE_URL  = f"{_BASE_URL}/api/erp/salesOrderLines/fieldChange?fieldName={{fieldName}}&dataOperation=CREATE"
SYNC_LINE_URL     = f"{_BASE_URL}/api/erp/salesOrderLinesGrid"

DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"]

# ── Mandatory columns — Header sheet ───────────────────────────────────────
MANDATORY_HEADER_COLUMNS = [
    "SO Number",
    "Domain Code",
    "Sold To Customer Code",
    "Site Code",
    "Currency Code",
    "Credit Terms",
    "Daybook Set",
]

# ── Mandatory columns — Lines sheet ─────────────────────────────────────────
MANDATORY_LINE_COLUMNS = [
    "SO Number",
    "Site Code",
    "Item Code",
    "Quantity Ordered",
]

# ── QAD API fieldName → Excel column name (Header / salesOrderHeaders) ─────
# Extend this map whenever a new QAD fieldName is observed in errors[]
# that should highlight a specific Excel cell. Anything NOT in this map
# still shows up in the Error message text — it just won't get its own
# highlighted cell.
SO_HEADER_FIELD_TO_COLUMN = {
    "salesOrderNumber":   "SO Number",
    "domainCode":         "Domain Code",
    "soldToCustomerCode": "Sold To Customer Code",
    "billToCustomerCode": "Bill To Customer Code",
    "shipToCustomerCode": "Ship To Customer Code",
    "siteCode":           "Site Code",
    "currencyCode":       "Currency Code",
    "daybookSetCode":     "Daybook Set",
    "creditTermsCode":    "Credit Terms",
    "shipVia":            "Ship Via",
    "freightListCode":    "Freight List",
    "freightTermsCode":   "Freight Terms",
    "orderDate":          "Order Date",
    "dueDate":            "Due Date",
}

# ── QAD API fieldName → Excel column name (Lines / salesOrderLines) ────────
SO_LINE_FIELD_TO_COLUMN = {
    "salesOrderNumber": "SO Number",
    "itemCode":         "Item Code",
    "siteCode":         "Site Code",
    "quantityOrdered":  "Quantity Ordered",
    "listPrice":        "List Price",
    "discountFormatted":"Discount",
    "netPrice":         "Net Price",
    "dueDate":          "Due Date",
}


# =============================================================================
# 1. AUTHENTICATION  (unchanged token pattern, now config-driven)
# =============================================================================

def _fetch_token() -> str:
    token_url = f"{_BASE_URL}/oauth/token"
    resp = requests.post(token_url, params=CONFIG["qad"]["auth"], timeout=30)
    resp.raise_for_status()
    token = resp.json().get("access_token")
    if not token:
        raise RuntimeError("OAuth response did not contain access_token")
    return token


class TokenManager:
    """Holds one token for the run; refreshes on demand (401)."""

    def __init__(self):
        self._token: str | None = None

    def get(self) -> str:
        if self._token is None:
            self._token = _fetch_token()
        return self._token

    def refresh(self) -> str:
        self._token = _fetch_token()
        return self._token

    def headers(self) -> dict:
        return {
            "Content-Type":  "application/json",
            "Authorization": f"Bearer {self.get()}",
        }


# =============================================================================
# 2. HTTP HELPERS  (auto-refresh on 401, one retry — UNCHANGED behaviour)
# =============================================================================

def api_get(url: str, tm: TokenManager) -> tuple[requests.Response, dict]:
    """GET with one automatic token-refresh retry on 401."""
    for attempt in range(2):
        resp = requests.get(url, headers=tm.headers(), timeout=30)
        if resp.status_code == 401 and attempt == 0:
            tm.refresh()
            continue
        try:
            return resp, resp.json()
        except Exception:
            return resp, {}
    return resp, {}


def api_post(url: str, payload: dict, tm: TokenManager) -> tuple[requests.Response, dict]:
    """POST with one automatic token-refresh retry on 401."""
    for attempt in range(2):
        resp = requests.post(url, json=payload, headers=tm.headers(), timeout=30)
        if resp.status_code == 401 and attempt == 0:
            tm.refresh()
            continue
        try:
            return resp, resp.json()
        except Exception:
            return resp, {}
    return resp, {}


# ── Normalized-error wrappers ────────────────────────────────────────────
# These wrap api_get()/api_post() so every caller (header creation, line
# creation) gets back the same (resp, resp_json, errors) shape used
# throughout Customer_load.py / PO_load.py — errors is either None
# (success) or a normalized list of {"fieldName", "message", "fieldValue",
# "code"} dicts, ready to be handed straight to resolve_qad_errors().

def _network_error(exc: Exception) -> dict:
    return {"fieldName": None, "message": f"Network error: {exc}", "fieldValue": None, "code": None}


def _token_error() -> dict:
    return {"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}


def _http_error(status_code: int) -> dict:
    return {"fieldName": None, "message": f"Unexpected response — HTTP {status_code}", "fieldValue": None, "code": None}


def safe_get(url: str, tm: TokenManager) -> tuple[requests.Response | None, dict | None, list | None]:
    """GET wrapped with network-error trapping. Does NOT run normalize_response
    (GET responses here are plain data payloads, not submit envelopes) —
    mirrors safe_get() in PO_load.py / get_customer() in Customer_load.py."""
    try:
        resp, resp_json = api_get(url, tm)
    except requests.RequestException as e:
        return None, None, [_network_error(e)]

    if resp.status_code == 401:
        return resp, resp_json, [_token_error()]
    if resp.status_code != 200:
        return resp, resp_json, [_http_error(resp.status_code)]
    return resp, resp_json, None


def safe_post(url: str, payload: dict, tm: TokenManager) -> tuple[requests.Response | None, dict | None, list | None]:
    """POST wrapped with network-error trapping and normalize_response(),
    exactly like safe_post() in PO_load.py / post_customer() in Customer_load.py."""
    try:
        resp, resp_json = api_post(url, payload, tm)
    except requests.RequestException as e:
        return None, None, [_network_error(e)]

    if resp.status_code == 401:
        return resp, resp_json, [_token_error()]

    success, errors = normalize_response(resp_json, resp.status_code)
    if not success:
        if not errors:
            errors = [_http_error(resp.status_code)]
        return resp, resp_json, errors
    return resp, resp_json, None


# =============================================================================
# 3. VALUE EXTRACTORS  (unchanged)
# =============================================================================

def sv(row_data: dict, key: str, default: str = "") -> str:
    val = row_data.get(key, default)
    return str(val).strip() if val is not None else default


def fv(row_data: dict, key: str, default: float = 0.0) -> float:
    try:
        return float(row_data.get(key, default))
    except (TypeError, ValueError):
        return default


def iv(row_data: dict, key: str, default: int = 0) -> int:
    try:
        return int(float(row_data.get(key, default)))
    except (TypeError, ValueError):
        return default


def to_date(val):
    """Convert a cell value (datetime, Excel serial number, or string in
    one of DATE_FORMATS) to an ISO 8601 UTC string. Unchanged from the
    original SO loader — kept because it handles Excel serial dates,
    which the simpler PO to_iso_date() does not need to."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%dT00:00:00.000Z")
    if isinstance(val, (int, float)):
        try:
            excel_epoch = datetime(1899, 12, 30)
            return (excel_epoch + timedelta(days=int(val))).strftime("%Y-%m-%dT00:00:00.000Z")
        except Exception:
            pass
    str_val = str(val).strip()
    if not str_val or str_val.lower() == "none":
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(str_val, fmt).strftime("%Y-%m-%dT00:00:00.000Z")
        except ValueError:
            continue
    return str_val


def today_iso():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT00:00:00.000Z")


def strip_keys(d: dict) -> dict:
    return {k.strip() if k else k: v for k, v in d.items()}


# =============================================================================
# 4. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict, columns: list[str]) -> list[str]:
    missing = []
    for col in columns:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 5. WORKBOOK HELPERS  (sheet/header layout UNCHANGED — row 1 = headers)
# =============================================================================

def ensure_columns(ws, *col_names: str) -> list:
    """Add any missing column names to row 1 and return the full header list."""
    header_row = [
        cell.value.strip() if isinstance(cell.value, str) else cell.value
        for cell in ws[1]
    ]
    for name in col_names:
        if name not in header_row:
            ws.cell(row=1, column=len(header_row) + 1, value=name)
            header_row.append(name)
    return header_row


# ── File rename helpers (same convention as PO_load.py) ────────────────────

def rename_error(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if not name.startswith("error_"):
        new_path = os.path.join(folder, "error_" + name)
        os.rename(file_path, new_path)
        return new_path
    return file_path


def rename_restore(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if name.startswith("error_"):
        new_path = os.path.join(folder, name[len("error_"):])
        os.rename(file_path, new_path)
        return new_path
    return file_path


# =============================================================================
# 6. PAYLOAD BUILDER  (UNCHANGED business logic)
# =============================================================================

def build_header_payload(row_data: dict) -> dict:
    domain   = sv(row_data, "Domain Code")
    so_num   = sv(row_data, "SO Number")
    order_dt = to_date(row_data.get("Order Date")) or today_iso()
    due_dt   = to_date(row_data.get("Due Date"))   or today_iso()
    uri      = f"urn:be:com.qad.sales.salesorder.ISalesOrderHeader:{domain}.{so_num}"

    sold_to       = sv(row_data, "Sold To Customer Code") or sv(row_data, "Bill To Customer Code")
    bill_to       = sv(row_data, "Bill To Customer Code")
    ship_to       = sv(row_data, "Ship To Customer Code")
    ship_via      = sv(row_data, "Ship Via")
    freight_list  = sv(row_data, "Freight List")
    freight_terms = sv(row_data, "Freight Terms")
    currency_code = sv(row_data, "Currency Code")
    site_code     = sv(row_data, "Site Code")
    credit_terms  = sv(row_data, "Credit Terms")
    daybook_set   = sv(row_data, "Daybook Set")

    return {
        "salesOrderHeaders": [
            {
                "uri":                    uri,
                "salesOrderNumber":       so_num,
                "domainCode":             domain,
                "soldToCustomerCode":     sold_to,
                "billToCustomerCode":     bill_to,
                "shipToCustomerCode":     ship_to,
                "siteCode":               site_code,
                "currencyCode":           currency_code,
                "daybookSetCode":         daybook_set,
                "creditTermsCode":        credit_terms,
                "shipVia":                ship_via,
                "freightListCode":        freight_list,
                "freightTermsCode":       freight_terms,
                "languageCode":           "us",
                "orderDate":              order_dt,
                "dueDate":                due_dt,
                "exchangeRate":           1,
                "exchangeRate2":          1,
                "exchangeRateType":       "",
                "isConfirmed":            True,
                "isFixedPrice":           True,
                "isTaxable":              False,
                "isPartialOK":            True,
                "isPrimarySO":            True,
                "isForecastConsumed":     True,
                "isReprice":              True,
                "isDisplayTaxAmounts":    True,
                "isPrintSalesOrder":      True,
                "isPrintPackList":        True,
                "isPrintInvoiceHistory":  True,
                "isProcessPostTrailer":   True,
                "isUsingConsignmentInventory": True,
                "maximumAgingDays":       90,
                "dataOperation":          "C",
                "concurrencyHash":        "",
                "disallowedActions":      "",
                "disallowedActionsMessage": "",
                "supplementaryMessages":  [],
                "SOSalespersons":         [],
                "SOSaveOptions":          [],
                "soDraftMappings":        [],
            }
        ]
    }


# =============================================================================
# 7. LINE CREATION FLOW  (API sequencing UNCHANGED — only response/error
#    handling now goes through safe_get()/safe_post() + normalize_response())
# =============================================================================

def create_line(
    domain:        str,
    so_num:        str,
    line_row_data: dict,
    line_number:   int,
    tm:            TokenManager,
) -> tuple[bool, list]:
    """
    Full stateful SO line creation:
      1. GET  initialize
      2. POST fieldChange(itemCode)
      3. POST fieldChange(siteCode)
      4. POST fieldChange(quantityOrdered)
      5. POST fieldChange(listPrice)         — only if > 0
      6. POST fieldChange(discountFormatted) — only if > 0
      7. POST fieldChange(netPrice)          — only if > 0
      8. POST fieldChange(dueDate)
      9. POST sync (salesOrderLinesGrid)

    Unlike the PO loader, the line number here comes from the Excel
    sheet (explicit "Line Number" column, or sequential position) and
    is applied to the initialized line — this is intentional SO
    business logic and is unchanged.

    Returns (success, errors), where errors is a normalized list of
    {"fieldName", "message", "fieldValue", "code"} dicts (empty on
    success) ready for resolve_qad_errors().
    """
    item_code = sv(line_row_data, "Item Code")
    site_code = sv(line_row_data, "Site Code")
    qty       = fv(line_row_data, "Quantity Ordered")
    price     = fv(line_row_data, "List Price")
    discount  = fv(line_row_data, "Discount", 0.0)
    net       = fv(line_row_data, "Net Price", 0.0)
    due_dt    = to_date(line_row_data.get("Due Date")) or today_iso()

    # ── 1. Initialize blank line ───────────────────────────────────────────
    resp, resp_json, errors = safe_get(INIT_LINE_URL.format(domain=domain, so=so_num), tm)
    if errors:
        return False, errors

    lines = (resp_json or {}).get("data", {}).get("salesOrderLines", [])
    if not lines:
        return False, [{"fieldName": None, "message": "Init returned no line object", "fieldValue": None, "code": None}]

    line = lines[0]
    line["salesOrderLine"] = line_number

    # ── Helper: run a fieldChange and chain the response ──────────────────
    def field_change(field_name, value):
        nonlocal line
        line[field_name] = value
        url = FIELD_CHANGE_URL.format(fieldName=field_name)
        r, rj, errs = safe_post(url, {"salesOrderLines": [line]}, tm)
        if errs:
            return False, errs
        updated = (rj or {}).get("data", {}).get("salesOrderLines", [])
        if not updated:
            return False, [{"fieldName": field_name, "message": f"fieldChange({field_name}) returned no line", "fieldValue": value, "code": None}]
        line = updated[0]
        return True, []

    # ── 2–8. Field changes in order ────────────────────────────────────────
    steps = [
        ("itemCode",        item_code),
        ("siteCode",        site_code),
        ("quantityOrdered", qty),
    ]
    if price > 0:
        steps.append(("listPrice", price))
    if discount > 0:
        steps.append(("discountFormatted", discount))
    if net > 0:
        steps.append(("netPrice", net))
    steps.append(("dueDate", due_dt))

    for field_name, value in steps:
        ok, errs = field_change(field_name, value)
        if not ok:
            return False, errs
        time.sleep(0.1)

    # ── 9. Sync / commit line ─────────────────────────────────────────────
    resp, resp_json, errors = safe_post(SYNC_LINE_URL, {"salesOrderLines": [line]}, tm)
    if errors:
        return False, errors

    return True, []


# =============================================================================
# 8. SINGLE-FILE PROCESSOR
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    print(f"\n{'─'*55}")
    print(f"  Loading: {os.path.basename(file_path)}")
    print(f"{'─'*55}")

    wb = openpyxl.load_workbook(file_path)

    if "Header" not in wb.sheetnames or "Lines" not in wb.sheetnames:
        print("ERROR: Workbook must have sheets named 'Header' and 'Lines'")
        raise RuntimeError("Workbook must have sheets named 'Header' and 'Lines'")

    ws_h = wb["Header"]
    ws_l = wb["Lines"]

    h_headers = ensure_columns(ws_h, "Status", "Error")
    l_headers = ensure_columns(ws_l, "Status", "Error")

    h_status_col = h_headers.index("Status") + 1
    h_error_col  = h_headers.index("Error")  + 1
    l_status_col = l_headers.index("Status") + 1
    l_error_col  = l_headers.index("Error")  + 1

    # ── Index header rows by SO number ────────────────────────────────────
    header_rows: dict[str, tuple[int, dict]] = {}
    for row_idx, row in enumerate(ws_h.iter_rows(min_row=2), start=2):
        row_values = [cell.value for cell in row]
        if not any(row_values):
            continue
        row_data = strip_keys(dict(zip(h_headers, row_values)))
        so_num   = sv(row_data, "SO Number")
        if so_num:
            header_rows[so_num] = (row_idx, row_data)

    # ── Index line rows by SO number ──────────────────────────────────────
    line_rows: dict[str, list[tuple[int, dict]]] = defaultdict(list)
    for row_idx, row in enumerate(ws_l.iter_rows(min_row=2), start=2):
        row_values = [cell.value for cell in row]
        if not any(row_values):
            continue
        row_data = strip_keys(dict(zip(l_headers, row_values)))
        so_num   = sv(row_data, "SO Number")
        if so_num:
            line_rows[so_num].append((row_idx, row_data))

    h_success = h_skip = h_fail = 0
    l_success = l_skip = l_fail = 0

    total_sos = len(header_rows)

    for idx, (so_num, (h_row_idx, h_row_data)) in enumerate(header_rows.items(), start=1):
        if _progress_callback:
            _progress_callback(idx, total_sos)

        print(f"\n  SO: {so_num}")

        domain = sv(h_row_data, "Domain Code")

        # ── Step 1: Create header ─────────────────────────────────────────
        if str(h_row_data.get("Status", "")).strip().upper() == "DONE":
            print("  Header — already DONE, skipping")
            h_skip += 1
        else:
            # Mandatory check before any API call
            missing = _check_mandatory(h_row_data, MANDATORY_HEADER_COLUMNS)
            if missing:
                h_fail += 1
                mark_error(
                    ws_h, h_row_idx, h_status_col, h_error_col, h_headers,
                    missing,
                    f"Missing mandatory fields: {', '.join(missing)}",
                )
                wb.save(file_path)
                continue

            payload = build_header_payload(h_row_data)
            resp, resp_json, errors = safe_post(HEADER_CREATE_URL, payload, tm)

            if errors:
                h_fail += 1
                bad_cols, error_msg = resolve_qad_errors(errors, SO_HEADER_FIELD_TO_COLUMN)
                mark_error(ws_h, h_row_idx, h_status_col, h_error_col, h_headers, bad_cols, error_msg)
                print(f"  Header FAILED: {so_num} — {error_msg}")
                wb.save(file_path)
                continue

            mark_done(ws_h, h_row_idx, h_status_col, h_error_col)
            h_success += 1
            print(f"  Header OK: {so_num}")
            wb.save(file_path)
            time.sleep(0.3)

        # ── Step 2: Create lines ───────────────────────────────────────────
        so_lines = line_rows.get(so_num, [])
        if not so_lines:
            print(f"  WARNING: No lines found for SO: {so_num}")
            continue

        for line_number, (l_row_idx, l_row_data) in enumerate(so_lines, start=1):
            if str(l_row_data.get("Status", "")).strip().upper() == "DONE":
                print(f"    Line {line_number} — already DONE, skipping")
                l_skip += 1
                continue

            missing = _check_mandatory(l_row_data, MANDATORY_LINE_COLUMNS)
            if missing:
                l_fail += 1
                mark_error(
                    ws_l, l_row_idx, l_status_col, l_error_col, l_headers,
                    missing,
                    f"Missing mandatory fields: {', '.join(missing)}",
                )
                wb.save(file_path)
                continue

            # SO-specific: explicit "Line Number" column overrides sequence,
            # unchanged from the original loader.
            explicit_line = iv(l_row_data, "Line Number", 0)
            line_no = explicit_line if explicit_line > 0 else line_number

            item = sv(l_row_data, "Item Code")
            qty  = fv(l_row_data, "Quantity Ordered")
            print(f"    → line {line_no}: {item} | qty={qty} | price={fv(l_row_data,'List Price')} "
                  f"disc={fv(l_row_data,'Discount')} net={fv(l_row_data,'Net Price')}")

            success, errors = create_line(domain, so_num, l_row_data, line_no, tm)

            if success:
                mark_done(ws_l, l_row_idx, l_status_col, l_error_col)
                l_success += 1
                print(f"    Line {line_no} OK: {item}")
            else:
                l_fail += 1
                bad_cols, error_msg = resolve_qad_errors(errors, SO_LINE_FIELD_TO_COLUMN)
                mark_error(ws_l, l_row_idx, l_status_col, l_error_col, l_headers, bad_cols, error_msg)
                print(f"    Line {line_no} FAILED: {item} — {error_msg}")

            wb.save(file_path)
            time.sleep(0.2)

    total_success = h_success + l_success
    total_fail    = h_fail    + l_fail

    print(f"\n  Headers — OK {h_success} | Skipped {h_skip} | Failed {h_fail}")
    print(f"  Lines   — OK {l_success} | Skipped {l_skip} | Failed {l_fail}")

    return total_success, total_fail


# =============================================================================
# 9. ORCHESTRATOR
# =============================================================================

def run(folder_path: str) -> tuple[int, int]:
    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    tm = TokenManager()

    total_success = 0
    total_fail    = 0

    for file_name in xlsx_files:
        file_path = os.path.join(folder, file_name)

        s, fail = process_file(file_path, tm)
        total_success += s
        total_fail    += fail

        if fail > 0:
            file_path = rename_error(file_path)
            print(f"\n  Errors found — renamed to: {os.path.basename(file_path)}")
        elif s > 0:
            file_path = rename_restore(file_path)

        print(f"\n  Load Summary {'-'*30}")
        print(f"    Rows loaded : {s}")
        print(f"    Rows failed : {fail}")
        if fail == 0:
            print("    Result      : ALL ROWS LOADED SUCCESSFULLY ✓")
        else:
            print("    Result      : COMPLETED WITH ERRORS — fix red rows and re-run")

    print(f"\n{'='*55}")
    print(f"  TOTAL — Success: {total_success} | Failed: {total_fail}")
    print(f"{'='*55}")

    return total_success, total_fail


# =============================================================================
# 10. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["sales_order"])
    )
    ok, fail = run(folder)
    sys.exit(0 if fail == 0 else 1)