"""
site_load.py
------------
Loads Site master records from all .xlsx files in the configured Site
folder into QAD via the sites API (urn:be:com.qad.base.item.ISite).

Sheet layout
------------
One flat sheet (row 1 = headers, row 2+ = data — same convention as
po_load.py / so_load.py, not Customer_load.py's two-row layout):

    Domain | Entity | Site | Description | Default inventory status |
    Data operation | Status | Error

CREATE only. The captured trace this loader is built from only showed a
POST /sites call — no GET-then-update flow was ever confirmed against a
live sandbox, so "Data operation" values other than C/blank are flagged
as a row error instead of silently being treated as a create (see the
"Data operation must be C" check in process_file()).

Auth
----
Same pattern as every other <Entity>_load.py in this tool: TokenManager
holds a session token + base_url handed in by the caller (main.py's
session, or the __main__ CLI login block below) — this script never
authenticates itself, since config.json has no stored service-account
credentials to authenticate with.

MINIMAL PAYLOAD — how it was derived
-------------------------------------------------------------------------
What was captured for /api/erp/sites was the RESPONSE of a successful
create (a full echoed record — every Site field QAD knows about,
including a long tail of custom*/siMstr*/transfer*/EMT* fields that were
never set by the caller and simply came back as "" / 0 / false defaults).
That response is NOT the request payload.

The payload below was reverse-engineered from it by keeping only the
fields that either (a) map directly to a column in Site_sample.xlsx, or
(b) are boilerplate every other build_*_payload() in this codebase sends
on a create (uri, concurrencyHash, disallowedActions,
disallowedActionsMessage, supplementaryMessages, dataOperation). Every
other field seen in the captured response (siteType, automaticLocations,
siMstr*, transfer*, custom*, EMTSupplier*, ...) is left out entirely and
allowed to take whatever default QAD assigns — this is the "send only
mandatory and important fields" minimal payload the response itself
doesn't tell us how to build, since it's an echo, not a spec.

ASSUMPTIONS (unconfirmed against a live sandbox — verify before trusting
in production, same as BOM_load.py's documented assumptions):
  * uri is built as f"urn:be:com.qad.base.item.ISite:{domain}.{site}",
    matching the {parent}.{code} convention used by every other loader's
    create uri (Customer_load's ICustomerV2, so_load's
    ISalesOrderHeader). The captured response's own uri
    ("urn:be:com.qad.base.item.ISite:10USA.") is missing its site-code
    segment entirely, which looks like an artifact of how that one test
    record happened to be created rather than the real convention, so it
    was not copied verbatim.
  * dataOperation: "C" — matches po_load.py / so_load.py / price_list_load.py's
    single-POST header creates. The captured response shows dataOperation
    as "" (blank), which is consistent with QAD simply not echoing it back
    either way, so it isn't evidence against "C".
"""

import os
import sys
import time
import requests
import openpyxl

# ── Progress callback (used by main.py SSE streaming) ──────────────────

_progress_callback = None

def set_progress_callback(callback):
    global _progress_callback
    _progress_callback = callback


# ── Path / config ─────────────────────────────────────────────────────────
ROOT_DIR    = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT_DIR)
sys.path.insert(0, SCRIPTS_DIR)
from config import CONFIG
from excel_format_utils import mark_done, mark_error, resolve_qad_errors
from qad_response_utils import normalize_response

# ── Mandatory columns ────────────────────────────────────────────────────
MANDATORY_COLUMNS = [
    "Domain",
    "Entity",
    "Site",
    "Description",
    "Default inventory status",
]

# ── QAD API fieldName → Excel column name ───────────────────────────────
# Extend whenever a new fieldName is observed in errors[]. Anything not in
# this map still shows up in the Error message text — it just won't get
# its own highlighted cell.
QAD_FIELD_TO_COLUMN = {
    "domainCode":             "Domain",
    "entityCode":              "Entity",
    "siteCode":                "Site",
    "description":             "Description",
    "defaultInventoryStatus":  "Default inventory status",
}


# =============================================================================
# 1. AUTHENTICATION
# =============================================================================

class TokenManager:
    """Holds the QAD access token + base_url for one run.

    token:    OAuth access token already obtained by main.py at login
              (stored server-side in SESSIONS[session_id]).
    base_url: the qracore base_url for the session's environment
              (same host used for the original OAuth call), e.g.
              CONFIG['environments']['TEST']['base_url'].
    """

    def __init__(self, token: str, base_url: str):
        if not token:
            raise ValueError("TokenManager requires a valid session access_token")
        if not base_url:
            raise ValueError("TokenManager requires a base_url")
        self._token    = token
        self.base_url  = base_url.rstrip("/")

    def get(self) -> str:
        return self._token

    def refresh(self) -> str:
        # No stored credentials to re-authenticate with — the session's
        # token came from the user's login and this script never had the
        # password.
        raise RuntimeError(
            "QAD session token expired mid-run and cannot be refreshed "
            "automatically — please log in again and re-run the load."
        )


class _TokenExpired(Exception):
    pass


# =============================================================================
# 2. SITE API
# =============================================================================

def post_site(payload: dict, token: str, base_url: str) -> tuple[bool, list]:
    """
    POST site payload to QAD. Returns (success, errors), normalized via
    qad_response_utils.normalize_response() so callers get a uniform list
    of dicts (fieldName/message/...) regardless of whether QAD sent back a
    submitResult-wrapped response or a bare top-level errors[].

    Raises _TokenExpired on 401.
    """
    url = f"{base_url}/api/erp/sites?viewUri=urn:be:com.qad.base.item.ISite"

    resp = requests.post(
        url,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type":  "application/json",
        },
        json=payload,
        timeout=30,
    )

    if resp.status_code == 401:
        raise _TokenExpired()

    try:
        resp_json = resp.json()
    except Exception:
        resp_json = {}

    return normalize_response(resp_json, resp.status_code)


# =============================================================================
# 3. VALUE EXTRACTORS
# =============================================================================

def sv(row_data: dict, key: str, default: str = "") -> str:
    val = row_data.get(key, default)
    return str(val).strip() if val is not None else default


# =============================================================================
# 4. MANDATORY FIELD CHECK
# =============================================================================

def _check_mandatory(row_data: dict) -> list[str]:
    missing = []
    for col in MANDATORY_COLUMNS:
        v = row_data.get(col)
        if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
            missing.append(col)
    return missing


# =============================================================================
# 5. WORKBOOK HELPERS  (row 1 = headers)
# =============================================================================

def ensure_columns(ws, *col_names: str) -> list:
    """Add any missing column names to row 1 and return the full header list.

    Stops at the first blank cell: real headers are always contiguous from
    column 1, so this naturally excludes blank_template.py's hidden
    dropdown-option columns further right (column 100+) from being read as
    if they were data columns. Without this, a dropdown's option list
    sharing row numbers with early data rows would make an otherwise-blank
    row look "non-empty" to process_file()'s any(row_values) check below.
    """
    header_row = []
    for cell in ws[1]:
        if cell.value is None:
            break
        header_row.append(cell.value.strip() if isinstance(cell.value, str) else cell.value)
    for name in col_names:
        if name not in header_row:
            ws.cell(row=1, column=len(header_row) + 1, value=name)
            header_row.append(name)
    return header_row


def rename_error(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if not name.startswith("error_"):
        new_path = os.path.join(folder, "error_" + name)
        os.rename(file_path, new_path)
        return new_path
    return file_path


def rename_restore(file_path: str) -> str:
    folder, name = os.path.dirname(file_path), os.path.basename(file_path)
    if name.startswith("error_"):
        new_path = os.path.join(folder, name[len("error_"):])
        os.rename(file_path, new_path)
        return new_path
    return file_path


# =============================================================================
# 6. PAYLOAD BUILDER  (minimal — see module docstring)
# =============================================================================

def build_payload(row_data: dict) -> dict:
    domain      = sv(row_data, "Domain")
    entity      = sv(row_data, "Entity")
    site_code   = sv(row_data, "Site")
    description = sv(row_data, "Description")
    inv_status  = sv(row_data, "Default inventory status")

    uri = f"urn:be:com.qad.base.item.ISite:{domain}.{site_code}"

    return {
        "supplementaryMessages": [],
        "sites": [{
            "uri":                     uri,
            "domainCode":              domain,
            "entityCode":              entity,
            "siteCode":                site_code,
            "description":             description,
            "defaultInventoryStatus":  inv_status,
            "dataOperation":           "C",
            "concurrencyHash":         "",
            "disallowedActions":       "",
            "disallowedActionsMessage": "",
            "supplementaryMessages":   [],
        }]
    }


# =============================================================================
# 7. SINGLE-FILE PROCESSOR
# =============================================================================

def process_file(file_path: str, tm: TokenManager) -> tuple[int, int]:
    print(f"\n{'─'*55}")
    print(f"  Loading: {os.path.basename(file_path)}")
    print(f"{'─'*55}")

    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    header_row = ensure_columns(ws, "Status", "Error")
    status_col = header_row.index("Status") + 1
    error_col  = header_row.index("Error")  + 1

    ok_count   = 0
    fail_count = 0

    total_rows = ws.max_row - 1

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if _progress_callback:
            _progress_callback(row_idx - 1, total_rows)

        row_values = [cell.value for cell in row]
        if not any(row_values[:len(header_row)]):
            continue

        row_data = dict(zip(header_row, row_values))

        status         = str(row_data.get("Status", "")).strip().upper()
        data_operation = str(row_data.get("Data operation", "")).strip().upper() or "C"

        if status == "DONE":
            print(f"  Row {row_idx} — already DONE, skipping")
            continue

        if data_operation != "C":
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                ["Data operation"],
                "Only Data operation = C (Create) is supported by this loader.",
            )
            wb.save(file_path)
            continue

        missing = _check_mandatory(row_data)
        if missing:
            fail_count += 1
            mark_error(
                ws, row_idx, status_col, error_col, header_row,
                missing,
                f"Missing mandatory fields: {', '.join(missing)}",
            )
            wb.save(file_path)
            continue

        payload = build_payload(row_data)
        site_code = sv(row_data, "Site")
        print(f"  Site: {site_code}")

        success = False
        errors: list = []

        for attempt in range(2):
            try:
                success, errors = post_site(payload, tm.get(), tm.base_url)
                break
            except _TokenExpired:
                if attempt == 0:
                    try:
                        tm.refresh()
                        continue
                    except RuntimeError as refresh_exc:
                        errors = [{"fieldName": None, "message": str(refresh_exc), "fieldValue": None, "code": None}]
                        break
                errors = [{"fieldName": None, "message": "Token refresh failed — unauthorised", "fieldValue": None, "code": None}]
                break
            except requests.RequestException as e:
                errors = [{"fieldName": None, "message": f"Network error: {e}", "fieldValue": None, "code": None}]
                break

        if not success:
            fail_count += 1
            bad_cols, error_msg = resolve_qad_errors(errors, QAD_FIELD_TO_COLUMN)
            mark_error(ws, row_idx, status_col, error_col, header_row, bad_cols, error_msg)
            print(f"  Site FAILED: {site_code} — {error_msg}")
            wb.save(file_path)
            time.sleep(0.1)
            continue

        mark_done(ws, row_idx, status_col, error_col)
        ok_count += 1
        print(f"  Site OK: {site_code}")
        wb.save(file_path)
        time.sleep(0.1)

    print(f"\n  Sites — OK {ok_count} | Failed {fail_count}")

    return ok_count, fail_count


# =============================================================================
# 8. ORCHESTRATOR
# =============================================================================

def run(folder_path: str, tm: TokenManager) -> tuple[int, int]:
    folder = os.path.abspath(folder_path)

    if not os.path.exists(folder):
        print(f"ERROR: Folder not found: {folder}")
        raise RuntimeError(f"Folder not found: {folder}")

    xlsx_files = [
        f for f in os.listdir(folder)
        if f.endswith(".xlsx") and not f.startswith("~$")
    ]

    if not xlsx_files:
        print(f"ERROR: No .xlsx files found in: {folder}")
        raise RuntimeError(f"No .xlsx files found in: {folder}")

    total_success = 0
    total_fail    = 0

    for file_name in xlsx_files:
        file_path = os.path.join(folder, file_name)

        s, fail = process_file(file_path, tm)
        total_success += s
        total_fail    += fail

        if fail > 0:
            file_path = rename_error(file_path)
            print(f"\n  Errors found — renamed to: {os.path.basename(file_path)}")
        elif s > 0:
            file_path = rename_restore(file_path)

        print(f"\n  Load Summary {'-'*30}")
        print(f"    Rows loaded : {s}")
        print(f"    Rows failed : {fail}")
        if fail == 0:
            print("    Result      : ALL ROWS LOADED SUCCESSFULLY ✓")
        else:
            print("    Result      : COMPLETED WITH ERRORS — fix red rows and re-run")

    print(f"\n{'='*55}")
    print(f"  TOTAL — Success: {total_success} | Failed: {total_fail}")
    print(f"{'='*55}")

    return total_success, total_fail


# =============================================================================
# 9. ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    import getpass
    import requests as _requests

    environment = os.environ.get("QAD_ENVIRONMENT", "TEST").upper()
    env_cfg     = CONFIG["environments"][environment]

    username = os.environ.get("QAD_USERNAME") or input("QAD username: ")
    password = os.environ.get("QAD_PASSWORD") or getpass.getpass("QAD password: ")

    token_resp = _requests.post(
        f"{env_cfg['base_url']}/oauth/token",
        data={
            "client_id":  env_cfg["client_id"],
            "username":   username,
            "password":   password,
            "grant_type": env_cfg.get("grant_type", "password"),
        },
        timeout=30,
    )
    token_resp.raise_for_status()
    access_token = token_resp.json().get("access_token")
    if not access_token:
        sys.exit("ERROR: OAuth response did not contain access_token")

    tm = TokenManager(token=access_token, base_url=env_cfg["base_url"])

    folder = os.path.abspath(
        os.path.join(ROOT_DIR, CONFIG["folders"]["site"])
    )
    ok, fail = run(folder, tm)
    sys.exit(0 if fail == 0 else 1)